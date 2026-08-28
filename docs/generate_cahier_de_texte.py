#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Cahier de texte du cours — un bloc par seance, a coller dans le logiciel
de scolarite (Ypareo ou autre).

    cp docs/acces.exemple.json docs/acces.json     # puis renseignez-le
    python3 docs/generate_cahier_de_texte.py

Produit docs/Cahier_de_texte.docx : une section par seance, separees par un
saut de page. Chacune se copie d'un bloc, sans retouche.

A PERSONNALISER : la liste SEANCES ne contient que deux exemples, qui
montrent la forme attendue.

Le document porte les identifiants du site de cours : ni lui ni
docs/acces.json n'entrent dans le depot, le .gitignore les ecarte.
"""

import io
import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ACCES = os.path.join(RACINE, "docs", "acces.json")
MODELE = os.path.join(RACINE, "docs", "acces.exemple.json")
SORTIE = os.path.join(RACINE, "docs", "Cahier_de_texte.docx")

BLEU = RGBColor(0x1B, 0x3A, 0x5B)
GRIS = RGBColor(0x5A, 0x66, 0x72)

# ─────────────────────────────────────────────────────────── les séances
#
# DEUX EXEMPLES, à remplacer par les vôtres — une entrée par séance, dans
# l'ordre. Chacune produit un bloc qui se copie d'un seul tenant dans le champ
# de votre logiciel de scolarité, sans retouche.
#
# La clé « repere » est celle qui compte : c'est la phrase que l'étudiant relira
# avant l'examen. Une séance sans repère n'a pas de point d'appui.
#
# Le gras s'écrit avec des **astérisques** dans les textes ci-dessous.
#
SEANCES = [
    {
        "titre": "Jour 1 — Matin · Titre de la séance",
        "objectif": "Ce que la séance doit permettre de faire à la fin. Un verbe "
                    "d'action observable, pas « sensibiliser » ni « aborder ».",
        "faire": [
            "Lire le **chapitre 1** et suivre ses ressources : capsules audio, "
            "schémas, auto-évaluations. Elles font partie du cours, pas du décor.",
            "**Une consigne opérationnelle**, avec sa date et son livrable s'il y "
            "en a un.",
            "Le polycopié de chaque chapitre est téléchargeable pour approfondir.",
        ],
        "repere": "La phrase à retenir de cette séance. Formulez-la comme un test : "
                  "« à la fin, vous devez savoir dire, devant tel cas, si… ».",
    },
    {
        "titre": "Jour 1 — Après-midi · Titre de la séance",
        "objectif": "L'objectif de la deuxième séance.",
        "faire": [
            "Lire le **chapitre 2**, avec ses ressources.",
            "**Constituer les équipes** si le cours en comporte, et communiquer "
            "leur composition au formateur.",
        ],
        "repere": "Le repère de cette séance-là.",
    },
]

# Les livrables notés du module. Le récapitulatif n'apparaît qu'une fois, à la
# dernière séance : répété partout, on cesse de le lire.
LIVRABLES = [
    ("Le premier livrable", "Date ou moment de remise", "individuel"),
    ("Le deuxième livrable", "Date ou moment de remise", "équipe"),
]

# ---------------------------------------------------------------- mise en forme
def gras(par, texte, taille=10.5, couleur=None, italique=False, bold=False):
    """Écrit `texte` en reconnaissant les **passages en gras**."""
    for i, bout in enumerate(texte.split("**")):
        if not bout:
            continue
        r = par.add_run(bout)
        r.font.size = Pt(taille)
        r.font.bold = bold or (i % 2 == 1)
        r.font.italic = italique
        if couleur is not None:
            r.font.color.rgb = couleur


def titre_seance(doc, texte):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(texte)
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = BLEU
    return p


def intertitre(doc, texte):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(texte.upper())
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = GRIS
    return p


def para(doc, texte, taille=10.5, italique=False, apres=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(apres)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    gras(p, texte, taille, italique=italique)
    return p


def puce(doc, texte):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    gras(p, texte)
    return p


def brut(par, texte, taille=10, bold=False):
    """Écrit le texte tel quel, sans interpréter les **.

    Indispensable pour le mot de passe : il contient lui-même des astérisques,
    que le parseur de gras avalerait. « ISM**2023 » sortait « ISM2023 », et
    personne n'aurait pu se connecter.
    """
    r = par.add_run(texte)
    r.font.size = Pt(taille)
    r.font.bold = bold
    return r


def bloc_acces(doc, acces):
    intertitre(doc, "Ressources pédagogiques")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    brut(p, "Site du cours : ")
    brut(p, acces["site"])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    brut(p, "Identifiant : ")
    brut(p, acces["login"], bold=True)
    brut(p, "   ·   Mot de passe : ")
    brut(p, acces["motdepasse"], bold=True)


# ---------------------------------------------------------------- document
def construire(acces):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    qcm = acces.get("qcm", "")
    courriel = acces.get("courriel") or "à l'adresse communiquée par le formateur"

    for i, s in enumerate(SEANCES):
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        titre_seance(doc, s["titre"])
        para(doc, "Intitulé du cours — parcours concernés",
             9.5, italique=True, apres=8)

        intertitre(doc, "Objectif de la séance")
        para(doc, s["objectif"])

        intertitre(doc, "Travail à faire")
        for f in s["faire"]:
            puce(doc, f % {"qcm": qcm} if "%(qcm)s" in f else f)

        intertitre(doc, "Repère")
        para(doc, s["repere"], italique=True)

        bloc_acces(doc, acces)

        # Le récapitulatif des livrables ne figure qu'une fois, à la dernière séance.
        if i == len(SEANCES) - 1:
            intertitre(doc, "Livrables notés du module")
            for nom, quand, qui in LIVRABLES:
                puce(doc, "**%s** — %s (%s)" % (nom, quand, qui))
            para(doc, "Tous les livrables sont à envoyer **par courriel**, %s, "
                      "à l'exception de la présentation orale qui se tient en direct." % courriel,
                 apres=4)
            para(doc, "**Étudiants absents le mercredi après-midi** : envoyez une **vidéo** de votre "
                      "présentation orale dans la semaine, et **au plus tard le 31 août**.", apres=4)
            para(doc, "Rappel : les livrables ne doivent porter **ni nom ni prénom** — ni dans le "
                      "texte, ni dans le nom de fichier, ni dans les métadonnées. Identifiez-vous "
                      "par votre numéro d'équipe.", italique=True)

    doc.save(SORTIE)
    return SORTIE


def main():
    if not os.path.isfile(ACCES):
        sys.exit("Fichier d'accès absent : %s\n"
                 "    cp %s %s" % (ACCES, MODELE, ACCES))
    with io.open(ACCES, encoding="utf-8") as f:
        acces = json.load(f)
    for cle in ("site", "login", "motdepasse", "qcm"):
        if not str(acces.get(cle, "")).strip() or "a-remplir" in str(acces.get(cle)):
            sys.exit("Fichier d'accès : « %s » n'est pas renseigné." % cle)
    chemin = construire(acces)
    print("  %s" % os.path.relpath(chemin, RACINE).replace(os.sep, "/"))
    print("  %d séances, une par page — à copier bloc par bloc dans votre logiciel de scolarité." % len(SEANCES))
    print("  Ce document porte le mot de passe du site : il est exclu du dépôt.")
    return 0


if __name__ == "__main__":
    sys.exit(main())

# -*- coding: utf-8 -*-
"""Les contrôles du polycopié.

  regles        dans le texte du Word : aucune adresse imprimée hors des mentions légales, aucune réponse d'activité
                (explications, corrections) ;
  lire_docx     le Word s'ouvre (python-docx) ; ses champs, comptés par genre ;
  champs_du_pdf les champs que Word n'a pas su calculer (« (p. ?) », « Signet non défini ») ;
  pdf           le nombre de pages, chaque page en PNG, les planches en vis-à-vis (pages paires à gauche) ;
  images_pdf    la résolution effective des images du PDF ;
  qr_du_pdf     chaque QR code imprimé, relu : l'adresse qu'il porte ;
  liens_des_qr  chaque QR code imprimé est un lien vers cette adresse, et sa légende aussi.
"""
import os
import re
import unicodedata
import zipfile

from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
ADRESSE = re.compile(r"https?://\S+|\bwww\.\S+", re.I)
# Une adresse sans protocole (« exemple.org ») : aucune ne s'imprime non plus.
DOMAINE = re.compile(r"\b(?:[a-z0-9-]+\.)+(?:com|org|net|io|fr|sh|dev|edu|live|app|eu)\b(?![\w-])")


def _norm(t):
    t = unicodedata.normalize("NFC", t or "").replace("’", "'").replace(" ", " ").replace(" ", " ")
    return re.sub(r"\s+", " ", t).strip().lower()


def textes_docx(chemin):
    """Les paragraphes du Word (corps, zones de texte, en-têtes et pieds), en texte brut, sans les champs : ni leur
    instruction, ni leur résultat provisoire (« ? », « mettre à jour les champs »)."""
    out = []
    with zipfile.ZipFile(chemin) as z:
        parties = [n for n in z.namelist() if re.match(r"word/(document|header\d*|footer\d*)\.xml$", n)]
        for nom in parties:
            racine = etree.fromstring(z.read(nom))
            for p in racine.iter(W + "p"):
                morceaux, dans_champ = [], 0
                for el in p.iter():
                    if el.tag == W + "fldChar":
                        t = el.get(W + "fldCharType")
                        dans_champ += 1 if t == "begin" else (-1 if t == "end" else 0)
                    elif el.tag == W + "t" and not dans_champ:
                        morceaux.append(el.text or "")
                texte = "".join(morceaux).strip()
                if texte:
                    out.append((nom, texte))
    return out


def regles(chemin, reponses=(), autorises=(), visibles="", produits=()):
    """Les écarts aux règles du livret dans le Word : liste de messages (vide si tout va bien). autorises : les
    débuts d'adresse permis (l'adresse du site, dans les mentions légales) ; visibles : le texte que les pages
    montrent elles-mêmes (lecture.texte_visible) — une phrase du chapitre que la correction d'une activité cite n'est
    pas une réponse imprimée ; produits : les noms de produits qui ont la forme d'une adresse."""
    erreurs = []
    paragraphes = textes_docx(chemin)
    tout = _norm(" ".join(t for _, t in paragraphes))
    visibles = _norm(visibles)
    for nom, t in paragraphes:
        permises = []
        for m in ADRESSE.finditer(t):
            if any(m.group(0).startswith(a) for a in autorises):
                permises.append((m.start(), m.end()))
            else:
                erreurs.append(f"adresse imprimée : « {m.group(0)} »")
        for m in DOMAINE.finditer(t):
            if any(a <= m.start() < b for a, b in permises) or m.group(0) in produits:
                continue
            erreurs.append(f"adresse imprimée : « {m.group(0)} » dans « {t[:90]} »")
    for r in reponses:
        cle = _norm(r)[:60]
        if len(cle) >= 25 and cle in tout and cle not in visibles:
            erreurs.append(f"réponse imprimée : « {r[:80]} »")
    return erreurs


def reponses_des_activites(activites):
    """Les textes qui donnent une réponse, dans les données des composants d'une page."""
    out = []
    for genre, d in activites.values():
        if genre == "checkpoint":
            for q in d.get("questions", []):
                out.append(q.get("explication", ""))
                out.extend(x for x in q.get("pourquoi", []) or [] if x)
        elif genre == "reveal":
            out.append(d.get("reponse", ""))
        elif genre == "spotError":
            out.extend(f.get("pourquoi", "") for f in d.get("fragments", []))
        elif genre == "branch":
            for e in d.get("etapes", []):
                out.extend(o.get("consequence", "") for o in (e.get("options") or e.get("choix") or []))
    return [re.sub(r"\*\*", "", x) for x in out if x]


def champs_du_pdf(chemin):
    """Les champs que Word n'a pas su calculer, relus dans le PDF : un signet introuvable, un renvoi resté « (p. ?) »,
    une table ou un index restés à « mettre à jour les champs ». Rend [(page, extrait)]."""
    import fitz
    motifs = re.compile(r"Erreur\s*!|Error!|Signet non défini|Bookmark not defined|\(p\.\s*\?\)|"
                        r"mettre à jour les champs|update the fields", re.I)
    out = []
    with fitz.open(chemin) as d:
        for i, page in enumerate(d, 1):
            texte = page.get_text()
            for m in motifs.finditer(texte):
                out.append((i, texte[max(0, m.start() - 40):m.end() + 20].replace("\n", " ")))
    return out


def lire_docx(chemin):
    """Le Word s'ouvre-t-il ? Ses paragraphes, tableaux, images, sections et champs (par genre)."""
    import docx
    d = docx.Document(chemin)
    champs = {}

    def compter(instruction):
        mot = instruction.strip().split(" ")[0].upper() if instruction.strip() else ""
        if mot:
            champs[mot] = champs.get(mot, 0) + 1

    with zipfile.ZipFile(chemin) as z:
        for nom in z.namelist():
            if not re.match(r"word/(document|header\d*|footer\d*)\.xml$", nom):
                continue
            racine = etree.fromstring(z.read(nom))
            pile = []
            for el in racine.iter(W + "fldChar", W + "instrText"):
                if el.tag == W + "fldChar":
                    genre = el.get(W + "fldCharType")
                    if genre == "begin":
                        pile.append([])
                    elif genre in ("separate", "end") and pile and pile[-1] is not None:
                        compter("".join(pile[-1]))
                        pile[-1] = None
                    if genre == "end" and pile:
                        pile.pop()
                elif pile and pile[-1] is not None:
                    pile[-1].append(el.text or "")
    images = len(d.element.body.xpath("//w:p//w:r/w:drawing/wp:inline"))
    return dict(paragraphes=len(d.paragraphs), tableaux=len(d.tables), images=images,
                sections=len(d.sections), champs=champs)


def pdf(chemin, dossier=None, dpi=110, planches=True):
    """Le nombre de pages du PDF ; avec dossier : chaque page en PNG (page-01.png…) et les planches en vis-à-vis
    (planche-01.png : la page 1 seule à droite, puis 2-3, 4-5…)."""
    import fitz
    with fitz.open(chemin) as d:
        n = d.page_count
        sorties = []
        if dossier:
            os.makedirs(dossier, exist_ok=True)
            for f in os.listdir(dossier):
                if re.match(r"(page|planche)-\d+\.png$", f):
                    os.remove(os.path.join(dossier, f))
            for i, page in enumerate(d, 1):
                sortie = os.path.join(dossier, f"page-{i:02d}.png")
                page.get_pixmap(dpi=dpi).save(sortie)
                sorties.append(sortie)
    if dossier and planches and sorties:
        _planches(sorties, dossier)
    return n, sorties


def _planches(pages, dossier):
    from PIL import Image
    ims = [Image.open(p) for p in pages]
    l, h = ims[0].size
    groupes = [(None, 0)] + [(i, i + 1 if i + 1 < len(ims) else None) for i in range(1, len(ims), 2)]
    for k, (g, d) in enumerate(groupes, 1):
        planche = Image.new("RGB", (2 * l + 24, h), (200, 205, 208))
        if g is not None:
            planche.paste(ims[g], (0, 0))
        if d is not None:
            planche.paste(ims[d], (l + 24, 0))
        planche.save(os.path.join(dossier, f"planche-{k:02d}.png"))
    for im in ims:
        im.close()


def images_pdf(chemin):
    """[(page, largeur px, largeur cm, ppp)] des images du PDF."""
    import fitz
    out = []
    with fitz.open(chemin) as d:
        for i, page in enumerate(d, 1):
            for info in page.get_image_info():
                x0, y0, x1, y1 = info["bbox"]
                largeur_cm = (x1 - x0) / 72 * 2.54
                if largeur_cm < 0.5:
                    continue
                out.append((i, info["width"], round(largeur_cm, 2), round(info["width"] / (largeur_cm / 2.54))))
    return out


def _matrices(urls):
    """{taille : {matrice : adresse}} : la matrice de chaque adresse (segno, même réglage que qr.py)."""
    import segno
    matrices = {}
    for url in set(urls):
        code = segno.make(url, error="m", micro=False, boost_error=False)
        m = tuple(tuple(bool(v) for v in rang) for rang in code.matrix)
        matrices.setdefault(len(m), {})[m] = url
    return matrices


def _lire_qr(info, matrices):
    """L'adresse que porte une image du PDF (d.extract_image) : None pour un QR code illisible ou inconnu, False pour
    une image qui n'est pas un QR code (pas carrée, pas en deux tons)."""
    import io as _io
    from PIL import Image
    l, h = info["width"], info["height"]
    if abs(l - h) > 2 or l < 60:
        return False
    im = Image.open(_io.BytesIO(info["image"])).convert("L")
    hist = im.histogram()
    if (sum(hist[:64]) + sum(hist[192:])) < 0.97 * l * h:
        return False
    noir = im.point(lambda v: 255 if v < 128 else 0)
    boite = noir.getbbox()
    if not boite:
        return False
    zone = im.crop(boite)
    lz, hz = zone.size
    for n, connues in matrices.items():
        cote = lz / n
        if abs(lz - hz) > cote or cote < 2:
            continue
        m = tuple(tuple(zone.getpixel((int((c + 0.5) * cote), int((r + 0.5) * cote))) < 128
                        for c in range(n)) for r in range(n))
        if m in connues:
            return connues[m]
    return None


def qr_du_pdf(chemin, urls):
    """Relit chaque QR code imprimé dans le PDF et dit quelle adresse il porte : chaque adresse posée donne sa matrice
    (segno) ; chaque image du PDF qui ressemble à un QR code (carrée, presque en deux tons) est seuillée, recadrée,
    échantillonnée au centre de chaque module, puis comparée aux matrices de même taille. Rend (trouvés, inconnus) :
    trouvés = [(page, adresse)], inconnus = [(page, taille en px)]."""
    import fitz
    matrices = _matrices(urls)
    trouves, inconnus = [], []
    with fitz.open(chemin) as d:
        vus = {}
        for i, page in enumerate(d, 1):
            for x in page.get_images(full=True):
                xref = x[0]
                if xref not in vus:
                    info = d.extract_image(xref)
                    vus[xref] = (info["width"], _lire_qr(info, matrices))
                largeur, resultat = vus[xref]
                if resultat is False:
                    continue
                if resultat:
                    trouves.append((i, resultat))
                else:
                    inconnus.append((i, largeur))
    return trouves, inconnus


def _aire(r):
    return 0.0 if r.is_empty else r.width * r.height


def liens_des_qr(chemin, urls):
    """Les liens des QR codes, relus dans le PDF. Chaque QR code imprimé doit être couvert par une annotation de lien
    vers l'adresse qu'il porte, et sa légende aussi : au moins une annotation de même adresse sur du texte, sous lui
    ou à côté de lui. Rend dict(qr=nombre de QR codes, liens={nature : nombre}, anomalies=[message])."""
    import fitz
    from urllib.parse import unquote
    matrices = _matrices(urls)
    bilan = dict(qr=0, liens=dict(qr=0, legendes=0, autres=0, internes=0), anomalies=[])
    with fitz.open(chemin) as d:
        lus = {}
        for i, page in enumerate(d, 1):
            externes, vus = [], set()
            for l in page.get_links():
                if l["kind"] == fitz.LINK_URI:
                    cle = (l.get("uri"), tuple(round(v, 1) for v in l["from"]))
                    if cle not in vus:
                        vus.add(cle)
                        externes.append((unquote(l.get("uri") or ""), fitz.Rect(l["from"])))
                elif l["kind"] in (fitz.LINK_GOTO, fitz.LINK_NAMED):
                    bilan["liens"]["internes"] += 1
            poses = []
            for x in page.get_image_info(xrefs=True):
                xref = x.get("xref") or 0
                if not xref:
                    continue
                if xref not in lus:
                    lus[xref] = _lire_qr(d.extract_image(xref), matrices)
                if lus[xref]:
                    poses.append((lus[xref], fitz.Rect(x["bbox"])))
            sur_qr = set()
            for url, r in sorted(poses, key=lambda x: (x[1].y0, x[1].x0)):
                bilan["qr"] += 1
                cible = unquote(url)
                image = {k for k, (u, z) in enumerate(externes) if u == cible and _aire(z & r) >= 0.8 * _aire(r)}
                autour = fitz.Rect(r.x0 - 330, r.y0 - 14, r.x1 + 330, r.y1 + 90)
                legende = [k for k, (u, z) in enumerate(externes)
                           if k not in image and u == cible and z.intersects(autour)
                           and _aire(z & r) < 0.5 * _aire(z) and page.get_textbox(z).strip()]
                sur_qr |= image
                if not image:
                    bilan["anomalies"].append(f"p. {i} : QR code sans lien vers son adresse ({url})")
                if not legende:
                    bilan["anomalies"].append(f"p. {i} : légende du QR code sans lien vers son adresse ({url})")
            adresses = {unquote(u) for u, _ in poses}
            for k, (u, z) in enumerate(externes):
                if k in sur_qr:
                    bilan["liens"]["qr"] += 1
                elif u in adresses:
                    bilan["liens"]["legendes"] += 1
                else:
                    bilan["liens"]["autres"] += 1
                    bilan["anomalies"].append(f"p. {i} : lien vers {u} sans QR code de même adresse sur la page")
    return bilan

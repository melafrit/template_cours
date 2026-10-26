# -*- coding: utf-8 -*-
"""Les pages propres au livret : celles que le site n'a pas.

  couverture           la page de garde : la composition du chemin du cours (pictos.py), le titre, le sous-titre,
                       l'auteur, le pied de couverture de la fiche du livret ;
  mentions             son verso : édition, copyright, usage, cas fictif (s'il y en a un), marques, le site du cours
                       et son QR code ;
  a_propos             « À propos de l'auteur » : le texte de la fiche du livret, mot pour mot ;
  mode_emploi          deux pages en vis-à-vis : la marge et ses QR codes, les réponses, la légende des encadrés, le
                       site ;
  table_des_matieres   le champ TOC (entrées de niveau 1 : champs TC ; niveau 2 : les sections) ;
  ouverture            l'ouverture d'un chapitre, sur une page de droite : numéro, titre, accroche, durée, objectifs,
                       mini-sommaire (champs PAGEREF vers les signets des sections) ; dans la marge, la page en ligne,
                       sa capsule et l'épisode du podcast du chapitre ;
  glossaire            les termes tels que les définit le glossaire du site, en deux colonnes ;
  index                le champ INDEX, alimenté par les champs XE que rendu.py pose dans le texte.

Chaque QR code et ce qui le décrit (sa légende ; à côté de lui, le titre, la phrase et l'adresse du site) mènent,
d'un clic dans le PDF, à son adresse (rendu.py, « LES LIENS »).
"""
import re
import unicodedata

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from . import gabarit as G
from . import lecture as H
from . import ooxml as X
from .rendu import signet_nom, vedette
from .sources import presentation_auteur

NBSP = " "


def _run(par, texte, police=None, taille=None, couleur=None, gras=False, italique=False, interlettre=None,
         capitales=False):
    r = par.add_run(texte)
    if police:
        G._police(r._r.get_or_add_rPr(), police)
    if taille:
        r.font.size = Pt(taille)
    if couleur:
        r.font.color.rgb = RGBColor.from_string(couleur)
    if gras:
        r.font.bold = True
    if italique:
        r.font.italic = True
    if capitales:
        r.font.all_caps = True
    if interlettre:
        X.poser(r._r.get_or_add_rPr(), X.xml(f'<w:spacing w:val="{int(interlettre * 20)}"/>'))
    return r


def _filet_accent(r, cont, longueur=3.0, largeur_texte=G.TEXTE, avant=6, apres=10, epaisseur=12):
    """Un court filet doré (la bordure basse d'un paragraphe vide, réduit par son retrait droit)."""
    p = r.nouveau_p(cont, "Poly Espace")
    p.paragraph_format.space_before = Pt(avant)
    p.paragraph_format.space_after = Pt(apres)
    p.paragraph_format.right_indent = Cm(max(0.0, largeur_texte - longueur))
    X.bordures_paragraphe(p, bottom=X.filet(sz=epaisseur, space=0, color=G.ACCENT))
    return p


def _pleine_page(r, hauteurs, largeur):
    """Un tableau sans bordure dont les lignes ont des hauteurs exactes : la page se compose à coup sûr."""
    t = r.doc.add_table(rows=len(hauteurs), cols=1)
    X.tableau_sans_bordures(t)
    X.tableau_largeurs(t, [largeur])
    for rang, h in zip(t.rows, hauteurs):
        X.hauteur_ligne(rang, h, "exact")
        X.ligne_insecable(rang)
        X.marges_cellule(rang.cells[0], 0, 0, 0, 0)
    return t


# ── la page de garde ───────────────────────────────────────────────────────
def couverture(r):
    doc, L, langue = r.doc, r.L, r.langue
    section = doc.sections[0]
    G.regler_section(section, "couverture", None, 1, premiere_page=True)
    G.entetes(section, doc, "couverture", langue)
    art = r.atelier.garde()
    ancre = r.nouveau_p(doc, "Poly Ancre")
    ancre.paragraph_format.space_before = Pt(0)
    ancre._p.append(X.xml(X.image_page_xml(r.rid(r.imprimable(art, G.PAGE_L)), G.PAGE_L, 13.5, 0, 0, r.ids(),
                                           L["alt"]["garde"], L["titre"] + " — " + L["sous_titre"])))
    t = _pleine_page(r, [11.55, 6.6, 3.9, 2.6], G.REFERENCE)
    c_titre, c_auteur, c_pied = t.rows[1].cells[0], t.rows[2].cells[0], t.rows[3].cells[0]
    p = r.nouveau_p(c_titre, "Poly Titre sans plan")
    p.paragraph_format.line_spacing = Pt(41)
    titre = L["titre"]
    coupe = titre.rfind(" ", 0, 28) if len(titre) > 28 else -1
    if coupe > 0:
        _run(p, r.typo(titre[:coupe]), G.TITRE, 34, G.PRIMAIRE_FONCE)
        p.add_run().add_break()
        _run(p, r.typo(titre[coupe + 1:]), G.TITRE, 34, G.PRIMAIRE_FONCE)
    else:
        _run(p, r.typo(titre), G.TITRE, 34, G.PRIMAIRE_FONCE)
    _filet_accent(r, c_titre, 3.2, G.REFERENCE, avant=10, apres=12, epaisseur=18)
    if L["sous_titre"]:
        s = r.nouveau_p(c_titre, "Normal")
        s.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.line_spacing = Pt(20)
        _run(s, r.typo(L["sous_titre"]), G.TEXTE_POLICE, 15, G.ARDOISE, italique=True)
    X.alignement_vertical(c_auteur, "bottom")
    a = r.nouveau_p(c_auteur, "Normal")
    a.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _run(a, L["auteur"], G.TITRE, 14, G.ENCRE, interlettre=0.6)
    X.alignement_vertical(c_pied, "bottom")
    f = r.nouveau_p(c_pied, "Normal")
    f.alignment = WD_ALIGN_PARAGRAPH.LEFT
    f.paragraph_format.space_after = Pt(0)
    X.bordures_paragraphe(f, top=X.filet(sz=6, space=8, color=G.ACCENT))
    _run(f, L["pied_couverture"] or NBSP, G.SANS, 10, G.PRIMAIRE, interlettre=1.6)


# ── le verso : mentions légales ────────────────────────────────────────────
def mentions(r):
    doc, L, M = r.doc, r.L, r.L["mentions"]
    G.nouvelle_section(doc, "couverture", r.langue, debut="suivante", premiere_page=True)
    t = _pleine_page(r, [G.HAUTEUR_UTILE - 0.35], G.REFERENCE)
    cell = t.rows[0].cells[0]
    X.alignement_vertical(cell, "bottom")
    mesure = 11.5

    def bloc(texte, police=G.TEXTE_POLICE, taille=8.5, couleur=G.ENCRE, apres=5, italique=False):
        p = r.nouveau_p(cell, "Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.right_indent = Cm(G.REFERENCE - mesure)
        p.paragraph_format.space_after = Pt(apres)
        p.paragraph_format.line_spacing = 1.18
        _run(p, r.typo(texte), police, taille, couleur, italique=italique)
        return p

    bloc(L["titre"], G.TITRE, 10, G.ENCRE, 0)
    if L["sous_titre"]:
        bloc(L["sous_titre"], G.TEXTE_POLICE, 9, G.ARDOISE, 0, italique=True)
    bloc(L["auteur"], G.SANS, 9, G.ENCRE, 16)
    bloc(M["edition"], G.SANS, 8.5, G.ENCRE, 3)
    bloc(M["copyright"], G.TITRE, 8.5, G.ENCRE, 8)
    bloc(M["usage"])
    if M.get("fiction"):
        bloc(M["fiction"])
    bloc(M["marques"], apres=16)
    # Le site du cours : son adresse (la seule imprimée du livret) et son QR code.
    url = r.langue.url("")
    png = r.registre.poser(url, M["site_titre"], "", "mentions", taille_cm=2.0)
    tq = cell.add_table(rows=1, cols=2)
    X.tableau_sans_bordures(tq)
    X.tableau_largeurs(tq, [2.4, mesure - 2.4])
    cq, ct = tq.rows[0].cells
    X.alignement_vertical(cq, "top")
    pq = r.nouveau_p(cq, "Poly Figure")
    pq.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pq.paragraph_format.space_before = Pt(0)
    pq._p.append(X.xml(r.qr_lie(png, 2.0, url, r.alt_qr(M["site_titre"], url))))
    lien = r.lien(url)
    p1 = r.nouveau_p(ct, "Poly Petit texte")
    p1.paragraph_format.space_after = Pt(2)
    _run(p1, M["site_titre"], G.TITRE, 8.5, G.PRIMAIRE)
    p2 = r.nouveau_p(ct, "Poly Petit texte")
    _run(p2, r.typo(M["site_texte"]), G.SANS, 8, G.ENCRE)
    p3 = r.nouveau_p(ct, "Poly Petit texte")
    _run(p3, url.rstrip("/"), G.SANS, 8, G.PRIMAIRE)
    for p in (p1, p2, p3):
        X.lier(p, lien)


# ── à propos de l'auteur ───────────────────────────────────────────────────
def a_propos(r):
    doc, L = r.doc, r.L
    texte = presentation_auteur(r.langue.code)
    saut = r.nouveau_p(doc, "Poly Ancre")
    saut.paragraph_format.page_break_before = True
    saut.paragraph_format.space_before = Pt(0)
    t = _pleine_page(r, [6.2, 12.0], G.REFERENCE)
    cell = t.rows[1].cells[0]
    marge_g = 2.2
    tp = r.nouveau_p(cell, "Poly Titre liminaire")
    tp.paragraph_format.left_indent = Cm(marge_g)
    tp.add_run(r.typo(L["a_propos"]))
    f = _filet_accent(r, cell, 2.6, G.REFERENCE - marge_g, avant=4, apres=14, epaisseur=14)
    f.paragraph_format.left_indent = Cm(marge_g)
    p = r.nouveau_p(cell, "Normal")
    p.paragraph_format.left_indent = Cm(marge_g)
    p.paragraph_format.right_indent = Cm(marge_g)
    p.paragraph_format.line_spacing = 1.42
    _run(p, r.typo(texte), G.TEXTE_POLICE, 11.5, G.ENCRE)
    r.journal.append(dict(genre="auteur", texte=texte))


# ── comment utiliser ce polycopié ──────────────────────────────────────────
def mode_emploi(r):
    doc, ME = r.doc, r.L["mode_emploi"]
    G.nouvelle_section(doc, "liminaire", r.langue, debut="paire", premiere_page=False)
    tp = r.nouveau_p(doc, "Poly Titre liminaire")
    tp.add_run(r.typo(ME["titre"]))
    X.ajouter_champ(tp, f"TC {X.guillemets_champ(r.typo(ME['titre']))} \\l 1", None)
    _filet_accent(r, doc, 3.0, avant=2, apres=12)
    p = r.nouveau_p(doc, "Normal")
    _run(p, r.typo(ME["intro"]))
    # La marge et ses QR codes.
    h = r.nouveau_p(doc, "Poly Sous-section")
    h.add_run(r.typo(ME["marge_titre"]))
    p = r.nouveau_p(doc, "Normal")
    _run(p, r.typo(ME["marge"]))
    r.marge(p, r.langue.url(""), r.L["marge"]["en_ligne"], r.L["marge"]["cours"], "ecran", place="mode d'emploi")
    t = doc.add_table(rows=len(ME["qr"]), cols=2)
    X.tableau_sans_bordures(t)
    X.tableau_largeurs(t, [3.6, G.TEXTE - 3.6])
    for rang, (picto, libelle, texte) in zip(t.rows, ME["qr"]):
        a, b = rang.cells
        for c in (a, b):
            X.bordures_cellule(c, bottom=X.filet(sz=4, color=G.FILET))
            X.marges_cellule(c, 0.09, 0.05, 0.09, 0.1)
        pa = r.nouveau_p(a, "Poly Étiquette")
        pa.paragraph_format.space_after = Pt(0)
        pa._p.append(r.run_picto(picto, G.PRIMAIRE, 0.4, -1.5))
        _run(pa, NBSP + NBSP + libelle, couleur=G.PRIMAIRE)
        pb = r.nouveau_p(b, "Poly Cellule")
        pb.paragraph_format.space_after = Pt(0)
        _run(pb, r.typo(texte[:1].upper() + texte[1:]), taille=9.5)
    r.nouveau_p(doc, "Poly Espace")
    h = r.nouveau_p(doc, "Poly Sous-section")
    h.add_run(r.typo(ME["reponses_titre"]))
    p = r.nouveau_p(doc, "Normal")
    _run(p, r.typo(ME["reponses"]))
    p = r.nouveau_p(doc, "Normal")
    _run(p, r.typo(ME["renvois"]))
    # Page de droite : les encadrés et le site.
    h = r.nouveau_p(doc, "Poly Sous-section")
    h.paragraph_format.page_break_before = True
    h.paragraph_format.space_before = Pt(0)
    h.add_run(r.typo(ME["encadres_titre"]))
    p = r.nouveau_p(doc, "Normal")
    _run(p, r.typo(ME["encadres"]))
    E = r.L["encadres"]
    for genre, texte in ME["legende"]:
        bx = r.boite(genre, doc)
        lp = r.etiquette(bx.entete, genre, E[genre])
        lp.paragraph_format.space_after = Pt(1)
        pc = r.nouveau_p(bx.entete, "Poly Petit texte")
        pc.paragraph_format.space_after = Pt(0)
        _run(pc, r.typo(texte[:1].upper() + texte[1:]), taille=9)
        bx.fermer()
    h = r.nouveau_p(doc, "Poly Sous-section")
    h.add_run(r.typo(ME["site_titre"]))
    p = r.nouveau_p(doc, "Normal")
    _run(p, r.typo(ME["site"]))


# ── la table des matières ──────────────────────────────────────────────────
def table_des_matieres(r):
    doc, L = r.doc, r.L
    tp = r.nouveau_p(doc, "Poly Titre sans plan")
    tp.paragraph_format.page_break_before = True
    tp.add_run(r.typo(L["toc"]))
    _filet_accent(r, doc, 3.0, avant=2, apres=16)
    p = r.nouveau_p(doc, "Normal")
    X.ajouter_champ(p, 'TOC \\o "2-2" \\u \\f \\l "1-2" \\h \\z', L["toc"] + " — " + L["maj_champs"])


# ── l'ouverture d'un chapitre ──────────────────────────────────────────────
def signet_page(rel):
    """Le signet d'une partie du livret (le titre de son ouverture) : la cible d'un renvoi vers la page entière."""
    return ("poly_page_" + re.sub(r"[^A-Za-z0-9]+", "_", rel.rsplit(".", 1)[0]))[:40]


def ouverture(r, page, n):
    """L'ouverture du chapitre n sur une page de droite : le numéro, le titre, l'accroche, la durée, les objectifs, le
    mini-sommaire (PAGEREF) ; dans la marge, la page en ligne, sa capsule et l'épisode du podcast du chapitre."""
    doc, L, O = r.doc, r.L, r.L["ouverture"]
    G.nouvelle_section(doc, "texte", r.langue, debut="impaire", premiere_page=True)
    r.page, r.chapitre = page, n
    esp = r.nouveau_p(doc, "Poly Espace")
    esp.paragraph_format.space_after = Pt(46)
    pn = r.nouveau_p(doc, "Poly Numéro de chapitre")
    _run(pn, L["chapitre"].format(n="").strip() + " ", G.TITRE, 11, G.ACCENT, interlettre=2.2, capitales=True)
    _run(pn, str(n), G.LEGER, 64, G.PRIMAIRE)
    pn.paragraph_format.line_spacing = Pt(70)
    pt = r.nouveau_p(doc, "Poly Titre de chapitre")
    pt.add_run(r.typo(page.titre))
    r.signet(pt, None, signet_page(page.rel))
    entree = L["chapitre_toc"].format(n=n, titre=r.typo(page.titre))
    X.ajouter_champ(pt, f"TC {X.guillemets_champ(entree)} \\l 1", None)
    r.marge(pt, r.langue.url(page.rel), r.L["marge"]["en_ligne"], r.L["marge"]["chapitre_en_ligne"].format(n=n),
            "ecran")
    r.podcast(pt, page.rel)
    _filet_accent(r, doc, 3.0, avant=4, apres=12, epaisseur=18)
    acc = r.nouveau_p(doc, "Normal")
    acc.alignment = WD_ALIGN_PARAGRAPH.LEFT
    acc.paragraph_format.line_spacing = 1.25
    acc.paragraph_format.space_after = Pt(8)
    r.ecrire(acc, page.sous_titre, base={"italique": True, "taille": 13, "couleur": G.ARDOISE}, indexer=False)
    if page.capsule_ouverture:
        r.capsule(page.capsule_ouverture, par=acc)
    if page.duree:
        pm = r.nouveau_p(doc, "Poly Petit texte")
        pm.paragraph_format.space_after = Pt(16)
        _run(pm, r.typo(O["duree"].format(duree=page.duree)), G.SANS, 8.5, G.ARDOISE)
    # Les objectifs.
    if page.objectifs:
        bx = r.boite("objectifs", doc)
        r.etiquette(bx.entete, "objectifs", page.objectifs_titre or r.L["encadres"]["objectifs"])
        cell = bx.ligne()
        for contenu in page.objectifs:
            pi = r.nouveau_p(cell, "Poly Encadré")
            G.Numerotation.appliquer(pi, r.num.puce)
            pi.paragraph_format.space_after = Pt(2)
            r.ecrire(pi, contenu, indexer=False)
        bx.fermer()
    # Le mini-sommaire : les sections et leur page (PAGEREF vers les signets posés sur les intertitres), et
    # « À retenir », dans l'ordre de la page.
    ps = r.nouveau_p(doc, "Poly Étiquette")
    ps.paragraph_format.space_before = Pt(10)
    ps.paragraph_format.space_after = Pt(6)
    ps._p.append(r.run_picto("liste", G.PRIMAIRE, 0.4, -1.6))
    _run(ps, NBSP + NBSP + O["sommaire"], couleur=G.PRIMAIRE)
    entrees, titres = [], dict(page.sections)
    for b in page.blocs:
        if isinstance(b, H.Titre) and b.niveau == 2 and b.ident in titres:
            entrees.append((b.ident, titres[b.ident]))
        elif isinstance(b, H.Encadre) and b.genre == "retenir" and page.retenir and b.etiquette == page.retenir:
            entrees.append((page.retenir, page.retenir_titre or O["retenir"]))
    for ident, titre in entrees:
        pe = r.nouveau_p(doc, "Poly Mini-sommaire")
        X.tabulation(pe, G.TEXTE, "right", points=True)
        m = re.match(r"(\d+)\.\s+(.*)", titre)
        if m:
            _run(pe, m.group(1) + ".", G.TITRE, couleur=G.ACCENT)
            pe.add_run(NBSP + NBSP + r.typo(m.group(2)))
        elif ident == page.retenir:
            _run(pe, r.typo(titre), G.TITRE, couleur=G.PRIMAIRE)
        else:
            _run(pe, "–", G.TITRE, couleur=G.ACCENT)
            pe.add_run(NBSP + NBSP + r.typo(titre))
        pe.add_run("\t")
        X.ajouter_champ(pe, f"PAGEREF {signet_nom(ident)} \\h", "?", X.rpr_xml(police=G.TITRE, couleur=G.PRIMAIRE))
        r.renvois.append(signet_nom(ident))
    r.saut_avant = True
    r.index_section, r.note_section = set(), False


# ── le glossaire ───────────────────────────────────────────────────────────
def _cle_tri(texte):
    s = unicodedata.normalize("NFD", texte.lower())
    return "".join(ch for ch in s if not unicodedata.combining(ch))


def _qr_bloc(r, cell, url, etiquette, legende, picto, taille):
    """Un QR code et sa légende (pictogramme et texte) dans une cellule ; l'un et l'autre mènent à son adresse."""
    png = r.registre.poser(url, etiquette, legende, "glossaire", taille_cm=taille)
    pq = r.nouveau_p(cell, "Poly Figure")
    pq.paragraph_format.space_before = Pt(0)
    pq._p.append(X.xml(r.qr_lie(png, taille, url, r.alt_qr(legende, url))))
    pl = r.nouveau_p(cell, "Poly Marge")
    pl.paragraph_format.space_before = Pt(2)
    lien = r.lien(url)
    pl._p.append(r.run_picto(picto, G.PRIMAIRE, 0.27, -1.4, lien=lien))
    _run(pl, NBSP + legende, taille=7)
    X.lier(pl, lien)


def glossaire(r, entrees, temoin=False, total=None):
    doc, Gl = r.doc, r.L["glossaire"]
    G.nouvelle_section(doc, "reference", r.langue, debut="impaire", premiere_page=True)
    r.page, r.chapitre = None, None
    tp = r.nouveau_p(doc, "Poly Titre de partie")
    tp.add_run(r.typo(Gl["titre"]))
    r.signet(tp, None, "poly_page_glossaire")
    X.ajouter_champ(tp, f"TC {X.guillemets_champ(r.typo(Gl['titre']))} \\l 1", None)
    _filet_accent(r, doc, 3.0, G.REFERENCE, avant=2, apres=8)
    ti = doc.add_table(rows=1, cols=2)
    X.tableau_sans_bordures(ti)
    X.tableau_largeurs(ti, [G.REFERENCE - 3.4, 3.4])
    ca, cb = ti.rows[0].cells
    pa = r.nouveau_p(ca, "Normal")
    pa.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pa.paragraph_format.right_indent = Cm(0.6)
    _run(pa, r.typo(Gl["intro"]), taille=10)
    if temoin:
        pt = r.nouveau_p(ca, "Poly Petit texte")
        _run(pt, r.typo(Gl["intro_temoin"].format(n=total)), G.SANS, 8.5, G.ARDOISE, italique=True)
    _qr_bloc(r, cb, r.langue.url("glossaire.html"), r.L["marge"]["en_ligne"], Gl["en_ligne"], "glossaire", 1.8)
    G.section_continue(doc, 2)
    lettre = None
    autre = "en" if r.langue.code == "fr" else "fr"
    for e in sorted(entrees, key=lambda e: _cle_tri(e["terme"])):
        initiale = _cle_tri(e["terme"])[:1].upper()
        if initiale != lettre:
            lettre = initiale
            pl = r.nouveau_p(doc, "Poly Lettre")
            pl.add_run(lettre)
            X.bordures_paragraphe(pl, bottom=X.filet(sz=4, space=2, color=G.FILET))
        p1 = r.nouveau_p(doc, "Poly Glossaire terme")
        p1.add_run(r.typo(vedette(e["terme"], r.L)))
        if e.get(autre):
            _run(p1, "  " + e[autre], G.SANS, 8, G.ARDOISE, italique=True)
        p2 = r.nouveau_p(doc, "Poly Glossaire définition")
        _run(p2, r.typo(e.get("definition", "")))
    # Une section continue à une colonne referme le glossaire : Word équilibre les deux colonnes de sa dernière page.
    G.section_continue(doc, 1)
    r.journal.append(dict(genre="glossaire", termes=len(entrees)))


# ── l'index ────────────────────────────────────────────────────────────────
def index(r):
    doc, I = r.doc, r.L["index"]
    G.nouvelle_section(doc, "reference", r.langue, debut="impaire", premiere_page=True)
    tp = r.nouveau_p(doc, "Poly Titre de partie")
    tp.add_run(r.typo(I["titre"]))
    X.ajouter_champ(tp, f"TC {X.guillemets_champ(r.typo(I['titre']))} \\l 1", None)
    _filet_accent(r, doc, 3.0, G.REFERENCE, avant=2, apres=8)
    pi = r.nouveau_p(doc, "Poly Petit texte")
    pi.paragraph_format.space_after = Pt(10)
    _run(pi, r.typo(I["intro"]), G.SANS, 8.5, G.ARDOISE, italique=True)
    p = r.nouveau_p(doc, "Normal")
    X.ajouter_champ(p, f'INDEX \\c "3" \\z "{r.L["lcid"]}" \\h "A"', I["titre"] + " — " + r.L["maj_champs"])

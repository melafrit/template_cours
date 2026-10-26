# -*- coding: utf-8 -*-
"""Les blocs d'une page → le document Word.

LA MARGE. Une zone de texte centrée dans la marge extérieure, ancrée sur la ligne qu'elle sert, porte un QR code
(2,2 cm), un libellé en capitales précédé de son pictogramme et une légende courte ; jamais d'adresse. On y trouve :
  - à l'ouverture de chaque chapitre, la page en ligne et, s'il en a un, l'épisode de son podcast (outils/podcasts.py,
    la page Podcasts à l'ancre de l'épisode) ;
  - la capsule audio de la section (l'adresse directe du MP3) ;
  - la définition d'un terme du glossaire, à sa première rencontre dans le livret (une par section) ;
  - le corrigé d'un point de contrôle, l'activité en ligne, la réponse d'une prédiction, le schéma interactif, les
    flashcards du chapitre ;
  - une ressource que le livret ne contient pas, désignée par un lien (une autre page du site, un autre site).
Une zone demandée dans un encadré s'accroche au paragraphe d'ancrage posé juste avant lui (Word placerait mal une
zone ancrée dans une cellule) ; deux zones d'un même paragraphe s'empilent ; word.py vérifie enfin, page par page,
qu'aucune ne déborde ni n'en chevauche une autre.

LES LIENS. Chaque QR code du livret et le texte qui le décrit — sous lui (la ligne du pictogramme et du libellé, la
légende, le terme et sa définition) ou à côté de lui — mènent, d'un clic dans le PDF, à l'adresse que porte le QR code
(qr_lie, ooxml.lier) ; l'impression ne change pas (ooxml, « LES LIENS »).

LES ENCADRÉS sont des tableaux d'une cellule, distincts même en noir et blanc :
  À retenir         filet épais bleu nuit en haut, fin en bas, fond bleuté très clair, une clé ;
  Repère            cadre bleu nuit fin, la boussole ;
  Attention, piège  filet épais brique à gauche, fond rosé, le triangle ;
  Activité          filet épais bleu nuit à gauche, fond gris clair, le crayon (point de contrôle : la case
                    cochée ; prédiction : la bulle) ;
  Pour aller plus loin  filets fins en haut et en bas, sans fond, les livres ;
  Lire ce schéma    filet bleu nuit à gauche, l'œil.

LES RÉPONSES ne s'impriment jamais : d'un point de contrôle, les questions et leurs propositions ; d'une association,
les deux colonnes, la droite mélangée ; d'un repérage d'erreurs, les passages et des cases vides ; d'une prédiction,
la question et des lignes pour écrire (ou ses propositions) ; d'un cas à embranchements, la situation et les choix.
Le reste est en ligne, derrière le QR code.

L'INDEX. Chaque terme du glossaire que le texte emploie reçoit un champ XE, une fois par section (les formes et les
exclusions de outils/glossaire_emplois.py, la règle de « Où en parle-t-on ? ») ; la page où la marge le définit est
en gras (\\b).

LES RENVOIS. Un lien vers une partie que le livret contient (une section d'un autre chapitre) devient « (p. n) », un
champ PAGEREF vers le signet de sa cible (Rendu.cibles) ; chaque signet visé doit être posé, le contrôle le vérifie.
AUCUNE ADRESSE ne s'imprime : un lien dont le texte est une adresse prend le nom que lui donne la fiche du livret
(« adresses » de outils/livret.py) et son QR code va dans la marge ; un code qui porte une adresse se copie en ligne.
"""
import os
import random
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import gabarit as G
from . import lecture as H
from . import ooxml as X
from .sources import GE, SITE_LOCAL

NBSP = " "
CASE = "☐"                       # ☐, en Segoe UI Symbol
SYMBOLE = "Segoe UI Symbol"
TELECHARGEABLES = (".pdf", ".docx", ".zip", ".xlsx", ".pptx")

# Pictogramme, couleur du filet et du libellé, fond, bordures de chaque genre d'encadré.
STYLES_ENCADRES = {
    "retenir": dict(picto="cle", couleur=G.PRIMAIRE, fond=G.FOND_PRIMAIRE,
                    bords=dict(top=X.filet(sz=18, color=G.PRIMAIRE), bottom=X.filet(sz=6, color=G.PRIMAIRE)),
                    marges=(0.22, 0.32, 0.16, 0.32)),
    "repere": dict(picto="boussole", couleur=G.PRIMAIRE, fond=G.BLANC,
                   bords={c: X.filet(sz=6, color=G.PRIMAIRE) for c in ("top", "left", "bottom", "right")},
                   marges=(0.18, 0.32, 0.14, 0.32)),
    "piege": dict(picto="alerte", couleur=G.BRIQUE, fond=G.FOND_PIEGE, bords=dict(left=X.filet(sz=24, color=G.BRIQUE)),
                  marges=(0.18, 0.35, 0.14, 0.3)),
    "activite": dict(picto="crayon", couleur=G.PRIMAIRE, fond=G.FOND_GRIS,
                     bords=dict(left=X.filet(sz=24, color=G.PRIMAIRE)), marges=(0.18, 0.35, 0.16, 0.3)),
    "controle": dict(picto="coche", couleur=G.PRIMAIRE, fond=G.FOND_GRIS,
                     bords=dict(left=X.filet(sz=24, color=G.PRIMAIRE)), marges=(0.18, 0.35, 0.16, 0.3)),
    "situation": dict(picto="question", couleur=G.PRIMAIRE, fond=G.FOND_GRIS,
                      bords=dict(left=X.filet(sz=24, color=G.PRIMAIRE)), marges=(0.18, 0.35, 0.16, 0.3)),
    "plus_loin": dict(picto="livres", couleur=G.PRIMAIRE, fond=None,
                      bords=dict(top=X.filet(sz=4, color=G.PRIMAIRE), bottom=X.filet(sz=4, color=G.PRIMAIRE)),
                      marges=(0.16, 0.05, 0.14, 0.05)),
    "lecture": dict(picto="oeil", couleur=G.PRIMAIRE, fond=None, bords=dict(left=X.filet(sz=10, color=G.PRIMAIRE)),
                    marges=(0.06, 0.32, 0.04, 0.0)),
    "objectifs": dict(picto="cible", couleur=G.PRIMAIRE, fond=G.FOND_PRIMAIRE, bords={},
                      marges=(0.2, 0.32, 0.16, 0.32)),
}
PICTOS_MARGE = ("casque", "micro", "glossaire", "crayon", "coche", "ecran", "question", "cartes", "oeil", "code")


class Glossaire:
    """Le glossaire du site et ses formes : où un texte emploie-t-il un terme ? Les formes et les expressions écartées
    sont celles de « Où en parle-t-on ? » (outils/glossaire_emplois.py)."""

    def __init__(self, langue):
        self.entrees = langue.glossaire()
        self.formes, self.exclus = {}, {}
        for e in self.entrees:
            for forme in GE.formes(e["terme"], e.get("alias"), langue.code):
                self.formes.setdefault(forme, e["ancre"])
            if e.get("exclus"):
                self.exclus[e["ancre"]] = GE.motif([GE.norm(x) for x in e["exclus"]])
        self.par_ancre = {e["ancre"]: e for e in self.entrees}
        self.motif = GE.motif(list(self.formes))

    def chercher(self, texte):
        """[(début, fin, entrée)] : les emplois de termes, sans chevauchement, la forme la plus longue d'abord."""
        if self.motif is None:
            return []
        norm = texte.lower().replace("’", "'").replace("‘", "'").replace(NBSP, " ").replace(" ", " ")
        out = []
        for m in self.motif.finditer(norm):
            ancre = self.formes[m.group(0)]
            ex = self.exclus.get(ancre)
            if ex and any(z.start() < m.end() and z.end() > m.start() for z in ex.finditer(norm)):
                continue
            out.append((m.start(), m.end(), self.par_ancre[ancre]))
        return out


class Rendu:
    def __init__(self, doc, langue, atelier, registre, glossaire=None):
        self.doc, self.langue, self.L = doc, langue, langue.L
        self.atelier, self.registre = atelier, registre
        self.ids = X.Compteur(1000)
        self.signets = X.Compteur(10)
        self.num = G.Numerotation(doc)
        self.glossaire = glossaire or Glossaire(langue)
        self.pictos = {}
        self.page = None                   # la page du site en cours (lecture.Page)
        self.chapitre = None               # son numéro
        self.index_section = set()         # les termes déjà indexés dans la section
        self.note_section = False          # la section a déjà sa définition dans la marge
        self.termes_notes = set()          # les termes déjà définis dans la marge du livret
        self.termes_vus = set()            # les termes que le livret emploie (le glossaire du témoin)
        self.dernier_p = None              # le dernier paragraphe du corps (ancre d'une capsule)
        self.ancre = None                  # le paragraphe d'ancrage de l'encadré en cours
        self.empile = {}                   # paragraphe → hauteur déjà occupée dans la marge (cm)
        self.cellules_entete = set()       # les en-têtes d'encadré : seuls paragraphes de cellule « liés au suivant »
        self.utilises = set()              # les paragraphes déjà écrits
        self.largeur = [G.TEXTE]           # la largeur disponible (cm), encadrés imbriqués compris
        self.style_texte = ["Normal"]
        self.indexer = [True]
        self.saut_avant = False            # le prochain paragraphe du corps commence une nouvelle page
        self.ancre_note = None             # l'élément qui termine le terme défini dans la marge (ecrire)
        self.journal = []
        self.cibles = {}                   # adresse du site (et ancre) → signet du livret : les renvois « p. n »
        self.signets_poses = set()
        self.renvois = []                  # les signets visés par un champ PAGEREF (contrôlés à la fin)
        self.qr_liens_vus = set()          # (page, adresse) des liens hors du livret déjà donnés en QR code
        self.podcasts = []                 # les épisodes du podcast posés aux ouvertures

    # ── outils ─────────────────────────────────────────────────────────────
    def preparer_pictos(self):
        demandes = {(n, G.PRIMAIRE) for n in PICTOS_MARGE}
        for st in STYLES_ENCADRES.values():
            demandes.add((st["picto"], st["couleur"]))
        demandes |= {("liste", G.PRIMAIRE), ("cible", G.PRIMAIRE), ("telephone", G.PRIMAIRE)}
        self.pictos = self.atelier.pictos(demandes)

    def picto(self, nom, couleur):
        cle = (nom, couleur)
        if cle not in self.pictos:
            self.pictos.update(self.atelier.pictos({cle}))
        return self.pictos[cle]

    def rid(self, chemin):
        return self.doc.part.get_or_add_image(chemin)[0]

    def imprimable(self, chemin, largeur_cm, dpi=255):
        """L'image ramenée à sa taille imprimée, à 255 points par pouce : l'export PDF de Word garde une image
        jusqu'à 260 points par pouce environ et ramène au-delà à 200 ; la capture de Chrome, plus fine, se réduit
        donc ici (Lanczos), à la bonne taille. Les QR codes, en deux couleurs, n'en ont pas besoin : Word les garde."""
        from PIL import Image
        cible = max(8, round(largeur_cm / 2.54 * dpi))
        with Image.open(chemin) as im:
            l, h = im.size
            if l <= cible * 1.03:
                return chemin
            base, ext = os.path.splitext(os.path.basename(chemin))
            dossier = os.path.join(os.path.dirname(os.path.dirname(chemin)), "imprimables")
            os.makedirs(dossier, exist_ok=True)
            sortie = os.path.join(dossier, f"{base}-{cible}{ext}")
            if not os.path.exists(sortie) or os.path.getmtime(sortie) < os.path.getmtime(chemin):
                petite = im.convert("RGBA" if im.mode in ("RGBA", "LA", "P") else "RGB").resize(
                    (cible, max(1, round(h * cible / l))), Image.LANCZOS)
                if ext.lower() in (".jpg", ".jpeg"):
                    petite.convert("RGB").save(sortie, "JPEG", quality=92, optimize=True)
                else:
                    petite.save(sortie, optimize=True)
        return sortie

    def run_picto(self, nom, couleur, taille=0.42, position=-1.5, lien=None):
        """Le run d'un pictogramme ; lien (r:id) : il est cliquable, comme la légende qu'il ouvre."""
        return X.xml(X.run_image_xml(self.rid(self.imprimable(self.picto(nom, couleur), taille)), taille, taille,
                                     self.ids(), self.L["alt"]["picto"], "", position, lien=lien))

    def alt_qr(self, legende, url):
        """Le texte de remplacement d'un QR code, dans la langue du livret : sa légende et son adresse."""
        return self.L["alt"]["qr"].format(legende=legende, url=url)

    def lien(self, url):
        """Le r:id de la relation « lien hypertexte » du document vers url : une seule par adresse."""
        return self.doc.part.relate_to(url, RT.HYPERLINK, is_external=True)

    def qr_lie(self, png, cote, url, alt):
        """Le XML d'un QR code cliquable (côté en cm) : l'image (a:hlinkClick), dans un w:hyperlink, mène à l'adresse
        qu'elle porte ; alt : son texte de remplacement."""
        lien = self.lien(url)
        return X.lien_xml([X.run_image_xml(self.rid(png), cote, cote, self.ids(), "QR code", alt, lien=lien)], lien)

    def nouveau_p(self, cont, style=None):
        """Un paragraphe dans le conteneur (document ou cellule) ; la cellule neuve offre son paragraphe vide."""
        if cont is not self.doc:
            ps = cont.paragraphs
            dernier = cont._tc[-1] if len(cont._tc) else None
            if (ps and dernier is ps[-1]._p and ps[-1]._p not in self.utilises and not ps[-1].text
                    and len(ps[-1]._p) <= 1):
                p = ps[-1]
                self.utilises.add(p._p)
                if style:
                    p.style = self.doc.styles[style]
                avant = p._p.getprevious()
                if avant is not None and avant.tag == qn("w:tbl"):
                    p.paragraph_format.space_before = Pt(5)     # de l'air sous un tableau imbriqué
                if getattr(cont, "_tc", None) not in self.cellules_entete:
                    p.paragraph_format.keep_with_next = False
                return p
        p = cont.add_paragraph(style=style)
        self.utilises.add(p._p)
        if cont is not self.doc and getattr(cont, "_tc", None) not in self.cellules_entete:
            p.paragraph_format.keep_with_next = False
        if self.saut_avant and cont is self.doc:
            p.paragraph_format.page_break_before = True
            self.saut_avant = False
        return p

    def tableau_cellule(self, cell, lignes, colonnes):
        """Un tableau imbriqué dans une cellule ; le paragraphe vide de la cellule neuve s'en va (il ferait un blanc
        au-dessus), celui que Word exige après le tableau servira au texte suivant."""
        ps = cell.paragraphs
        if len(ps) == 1 and len(cell._tc) and cell._tc[-1] is ps[0]._p and not ps[0].text and len(ps[0]._p) <= 1 \
                and len(cell._tc.findall(qn("w:tbl"))) == 0:
            cell._tc.remove(ps[0]._p)
        return cell.add_table(rows=lignes, cols=colonnes)

    def fin_de_cellule(self, cell):
        """Le paragraphe vide qui termine une cellule (après un tableau imbriqué) : une ligne d'un point."""
        ps = cell.paragraphs
        if ps and len(cell._tc) and cell._tc[-1] is ps[-1]._p and ps[-1]._p not in self.utilises \
                and not ps[-1].text and len(ps[-1]._p) <= 1 and cell._tc.findall(qn("w:tbl")):
            ps[-1].style = self.doc.styles["Poly Espace"]
            ps[-1].paragraph_format.space_after = Pt(0)

    def typo(self, texte):
        """La typographie du livret : apostrophe courbe, espaces insécables du français (une substitution d'un
        caractère par un autre : les positions ne bougent pas)."""
        t = texte.replace("'", "’")
        t = re.sub(r"(^|[\s(\[“«])’(?=\w)", r"\1‘", t)
        # Les guillemets droits deviennent typographiques : ouvrant en début de mot, fermant ailleurs.
        t = re.sub(r'(^|[\s(\[—–-])"', r"\1“", t).replace('"', "”")
        t = re.sub(r"§ (?=\d)", "§" + NBSP, t)
        if self.langue.code == "fr":
            t = re.sub(r" (?=[;:!?»])", NBSP, t)
            t = t.replace("« ", "«" + NBSP)
            t = re.sub(r"(?<=\d) (?=\d{3}\b)", NBSP, t)
            t = re.sub(r" (?=—[,.;:)])", NBSP, t)
        return t

    def signet(self, paragraphe, ident, nom=None):
        """Un signet sur un paragraphe (cible d'un PAGEREF) ; un nom déjà posé ne l'est pas deux fois."""
        nom = nom or signet_nom(ident)
        if nom in self.signets_poses:
            return nom
        X.signet(paragraphe, nom, self.signets())
        self.signets_poses.add(nom)
        return nom

    def signer(self, ident):
        """Un signet sur l'ancre de l'encadré en cours (une activité), pour les renvois de pages."""
        if ident and self.ancre is not None:
            self.signet(self.ancre, ident)

    def cible(self, url):
        """Le signet du livret que vise une adresse du site (avec son ancre, sinon la page), ou None."""
        base, _, ancre = url.partition("#")
        if ancre and (base, ancre) in self.cibles:
            return self.cibles[(base, ancre)]
        return self.cibles.get(base)

    def renvoi_page(self, par, nom, rpr="", entre_parentheses=True):
        """« (p. 12) » : un champ PAGEREF vers un signet du livret, mis à jour par Word."""
        avant, _, apres = self.L["renvoi_page"].partition("{n}")
        ouvre, ferme = (" (", ")") if entre_parentheses else ("", "")
        par._p.append(X.xml(X.run_xml(ouvre + avant, rpr)))
        for r in X.champ_xml(f"PAGEREF {nom} \\h", "?", rpr):
            par._p.append(X.xml(r))
        par._p.append(X.xml(X.run_xml(apres + ferme, rpr)))
        self.renvois.append(nom)

    def sans_adresses(self, contenu):
        """Aucune adresse imprimée : le lien dont le texte est une adresse prend le nom que lui donne la fiche du
        livret (« adresses ») ; rend (contenu, [(adresse, légende du QR code)])."""
        A = self.L.get("adresses", {})
        texte = "".join(s.texte for s in contenu)
        if not A or not any(dom in texte for dom in A):
            return contenu, []
        adresses = []
        for s in contenu:
            t = s.texte.strip().rstrip("/")
            for prefixe in ("https://", "http://"):
                if t.startswith(prefixe):
                    t = t[len(prefixe):]
            if s.lien and t in A:
                url = self.langue.resoudre(s.lien, self.page.rel) if self.page else s.lien
                if url not in [u for u, _ in adresses]:
                    adresses.append((url, A[t][1]))
        regles = [(re.escape(d), nom) for d, (nom, _) in A.items()]
        return reecrire(contenu, regles), adresses

    # ── le texte en ligne ──────────────────────────────────────────────────
    def hors_livret(self, url):
        """Un lien vers une ressource que le livret ne contient pas (une page du site absente du livret, un site
        extérieur), à donner en QR code : jamais une correction protégée."""
        return bool(re.match(r"https?://", url)) and "/correction/" not in url and not self.cible(url)

    def ecrire(self, par, contenu, base=None, indexer=True, noter=False, liens=True, refs=True):
        """Écrit un texte en ligne dans un paragraphe : runs, champs XE, liens en marge, renvois de pages. Rend
        (entrée du glossaire à définir dans la marge ou None, [(adresse, texte, élément) des liens à mettre en
        marge])."""
        base = base or {}
        contenu, adresses = self.sans_adresses(contenu)
        segs = [H.Seg(s.texte, s.gras, s.italique, s.code, s.lien, s.exposant) for s in contenu]
        brut = "".join(s.texte for s in segs)
        propre = self.typo(brut)
        i = 0
        for s in segs:
            n = len(s.texte)
            if not s.code:
                s.texte = propre[i:i + n]
            i += n
        # Les liens : vers une partie du livret → un renvoi de page ; vers un document ou une ressource hors du
        # livret → dans la marge.
        en_marge, par_seg, renvois, vus = [], {}, {}, set()
        for k, s in enumerate(segs):
            if not s.lien:
                continue
            url = self.langue.resoudre(s.lien, self.page.rel) if self.page else s.lien
            fleche = s.texte.rstrip().endswith("→")
            cible = self.cible(url) if refs else None
            if cible:
                if fleche:
                    s.texte = s.texte.replace(" →", "").replace("→", "")
                    s.italique = True
                propre_lien = s.texte.strip()
                if cible not in vus and not re.fullmatch(r"[\d\s.,]*", propre_lien):
                    renvois[k] = cible
                    vus.add(cible)
                continue
            if not liens:
                continue
            special = fleche or url.split("#")[0].lower().endswith(TELECHARGEABLES)
            if not (special or self.hors_livret(url)):
                continue
            cle = (self.page.rel if self.page else "", url.split("#")[0])
            if cle in self.qr_liens_vus:
                continue
            self.qr_liens_vus.add(cle)
            par_seg[k] = len(en_marge)
            en_marge.append([url, capitale(s.texte.replace("→", "").strip()), None])
            if fleche:
                s.texte = s.texte.replace(" →", "").replace("→", "")
                s.italique = True
        texte = "".join(s.texte for s in segs)
        # Les termes du glossaire : XE une fois par section ; la définition dans la marge à la première rencontre
        # du terme dans le livret, une par section.
        coupes, note = [], None
        if indexer and self.indexer[-1]:
            for debut, fin, e in self.glossaire.chercher(texte):
                terme = e["terme"]
                self.termes_vus.add(e["ancre"])
                gras = False
                if noter and note is None and not self.note_section and self.chapitre and terme not in self.termes_notes:
                    note, gras = e, True
                    self.note_section = True
                    self.termes_notes.add(terme)
                if terme in self.index_section and not gras:
                    continue
                self.index_section.add(terme)
                coupes.append((fin, terme, gras))
        # Les runs ; on retient l'élément qui termine le terme défini et chaque lien mis en marge : la zone de la
        # marge s'ancre juste après, sur sa ligne.
        pos, self.ancre_note = 0, None
        for k, s in enumerate(segs):
            debut, fin = pos, pos + len(s.texte)
            pos = fin
            ici = sorted((c for c in coupes if debut < c[0] <= fin), key=lambda c: c[0])
            curseur = debut
            for c in ici:
                self._run(par, s, s.texte[curseur - debut:c[0] - debut], base)
                instr = f"XE {X.guillemets_champ(self.typo(vedette(c[1], self.L)))}" + (" \\b" if c[2] else "")
                for r in X.champ_xml(instr, None):
                    par._p.append(X.xml(r))
                if c[2]:
                    self.ancre_note = par._p[-1]
                curseur = c[0]
            self._run(par, s, s.texte[curseur - debut:], base)
            if k in par_seg and len(par._p):
                en_marge[par_seg[k]][2] = par._p[-1]
            if k in renvois:
                rpr = X.rpr_xml(taille=base.get("taille"), couleur=base.get("couleur"), italique=base.get("italique"))
                self.renvoi_page(par, renvois[k], rpr)
        en_marge.extend([url, legende, None] for url, legende in adresses)
        return note, en_marge

    def _run(self, par, s, texte, base):
        if not texte:
            return
        if texte == "\n":
            par.add_run().add_break()
            return
        for k, m in enumerate(texte.split("\n")):
            if k:
                par.add_run().add_break()
            if not m:
                continue
            if not s.code:
                m = m.replace("—,", "—‍,").replace("—)", "—‍)")
            r = par.add_run(m)
            f = r.font
            if s.gras or base.get("gras"):
                f.bold = True
            if s.italique or base.get("italique"):
                f.italic = True
            if s.exposant:
                f.superscript = True
            if s.code:
                G._police(r._r.get_or_add_rPr(), G.CODE)
                f.size = Pt(base.get("taille", 10) - 1)
            elif base.get("police"):
                G._police(r._r.get_or_add_rPr(), base["police"])
            if base.get("couleur") and not s.code:
                f.color.rgb = RGBColor.from_string(base["couleur"])
            if base.get("taille") and not s.code:
                f.size = Pt(base["taille"])

    # ── la marge ───────────────────────────────────────────────────────────
    def lignes_zone(self, url, etiquette, texte="", picto="ecran", note=None, depart=0.0):
        """Ce qu'une zone de la marge porte sous son QR code : la ligne du pictogramme et du libellé, la légende, le
        terme et sa définition (note). Rend (le XML de ces paragraphes, depart augmenté de leur hauteur estimée, en
        cm). Tout y mène à url, l'adresse du QR code."""
        lien = self.lien(url)
        r_picto = X.run_image_xml(self.rid(self.imprimable(self.picto(picto, G.PRIMAIRE), 0.27)), 0.27, 0.27,
                                  self.ids(), self.L["alt"]["picto"], "", -1.4, lien=lien)
        contenu = [X.para_xml([X.lien_xml([r_picto, X.run_xml(NBSP + etiquette, X.rpr_xml(
            police=G.TITRE, taille=6.5, couleur=G.PRIMAIRE, espacement=0.35))], lien)], jc="center", interligne=9)]
        hauteur = depart + 0.09 + 0.32
        if texte:
            contenu.append(X.para_xml([X.lien_xml([X.run_xml(self.typo(texte), X.rpr_xml(
                police=G.SANS, taille=7, couleur=G.ARDOISE))], lien)], jc="center", interligne=8.6))
            hauteur += 0.31 * max(1, -(-len(texte) // 27))
        if note:
            terme, definition = note
            contenu.append(X.para_xml([X.lien_xml([X.run_xml(self.typo(terme), X.rpr_xml(
                police=G.TITRE, taille=7.5, couleur=G.ENCRE))], lien)], jc="left", avant=4, interligne=9.5))
            contenu.append(X.para_xml([X.lien_xml([X.run_xml(self.typo(definition), X.rpr_xml(
                police=G.SANS, taille=7, couleur=G.ENCRE))], lien)], jc="left", interligne=8.6))
            hauteur += 0.14 + 0.34 + 0.305 * max(1, -(-len(definition) // 30))
        return contenu, hauteur

    def marge(self, par, url, etiquette, texte="", picto="ecran", note=None, place="", apres=None):
        """Pose une zone dans la marge extérieure (voir l'en-tête du module) : ancrée juste après l'élément apres
        (la fin d'un terme, un lien), sinon au début du paragraphe par, où plusieurs zones s'empilent."""
        png = self.registre.poser(url, etiquette, texte, place or self.lieu())
        contenu = [X.para_xml([self.qr_lie(png, G.QR, url, self.alt_qr(f"{etiquette} {texte}".strip(), url))],
                              jc="center", apres=2.5)]
        lignes, hauteur = self.lignes_zone(url, etiquette, texte, picto, note, depart=G.QR)
        contenu += lignes
        decalage = 0.0 if apres is not None else self.empile.get(par._p, 0.0)
        if apres is None:
            self.empile[par._p] = decalage + hauteur + 0.35
        ident = self.ids()
        nom = f"Marge {ident} — {etiquette}"
        dessin = X.xml(X.zone_marge_xml(ident, nom, f"{etiquette} {texte}".strip(), G.MARGE_BOITE, hauteur, contenu,
                                        decalage))
        pp = par._p.find(qn("w:pPr"))
        if apres is not None and apres.getparent() is par._p:
            apres.addnext(dessin)
        elif pp is not None:
            pp.addnext(dessin)
        else:
            par._p.insert(0, dessin)
        self.journal.append(dict(genre="marge", etiquette=etiquette, texte=texte, url=url, lieu=self.lieu()))
        return dessin

    def poser_marges(self, par, corps, note, liens):
        """Les zones qu'un paragraphe demande : la définition du terme qu'il introduit, ses liens en marge. Dans le
        corps, chacune s'ancre sur sa ligne ; dans un encadré, à l'ancre de l'encadré."""
        cible = par if corps else self.cible_marge()
        if cible is None:
            return
        if note:
            self.marge(cible, self.langue.url("glossaire.html") + "#" + note["ancre"],
                       self.L["marge"]["glossaire"], "", "glossaire",
                       note=(vedette(note["terme"], self.L), note.get("definition", "")),
                       apres=self.ancre_note if corps else None)
        for url, texte, element in liens:
            self.marge(cible, url, self.L["marge"]["en_ligne"], texte, "ecran", apres=element if corps else None)

    def cible_marge(self):
        """Le paragraphe qui portera une zone de la marge : l'ancre de l'encadré en cours, sinon le dernier
        paragraphe du corps."""
        return self.ancre if self.ancre is not None else self.dernier_p

    def lieu(self):
        return self.page.rel if self.page is not None else ""

    # ── le podcast (module 03) ─────────────────────────────────────────────
    def episode_podcast(self, page_rel):
        """(adresse, légende) du QR code de l'épisode qui accompagne la page (outils/podcasts.py, clé « page »), ou
        None : la page Podcasts de la langue du livret, à l'ancre de l'épisode ; « Épisode 1 — Chapitre 1 »."""
        try:
            import podcasts as P
        except ImportError:                       # un cours sans série de podcasts
            return None
        ep = next((e for e in P.EPISODES if e.get("page") == page_rel), None)
        if ep is None:
            return None
        lg = self.langue.code
        legende = P.libelle("lecteur", lg, n=ep["num"], titre=P.texte(ep, "court", lg))
        return self.langue.url(f"podcasts.html#episode-{ep['num']}"), legende, ep["num"]

    def podcast(self, par, page_rel):
        """La zone PODCAST de l'ouverture d'un chapitre, s'il a un épisode : ancrée au paragraphe par (le titre de
        l'ouverture), sous les zones qu'il porte déjà."""
        episode = self.episode_podcast(page_rel)
        if not episode:
            return None
        url, legende, num = episode
        zone = self.marge(par, url, self.L["marge"]["podcast"], legende, "micro")
        # Dans le Word et le PDF balisé, la zone se lit après celles qu'elle suit (marge l'a mise en tête).
        suivant = zone.getnext()
        while suivant is not None and suivant.tag == qn("w:r") and suivant.find(".//" + qn("wp:anchor")) is not None:
            suivant.addnext(zone)
            suivant = zone.getnext()
        self.registre.poses[-1]["episode"] = num
        self.podcasts.append(num)
        return zone

    # ── les blocs ──────────────────────────────────────────────────────────
    def blocs(self, blocs, cont=None):
        cont = cont or self.doc
        for b in blocs:
            self.bloc(b, cont)

    def bloc(self, b, cont):
        if isinstance(b, H.Titre):
            return self.titre(b, cont)
        if isinstance(b, H.Paragraphe):
            return self.paragraphe(b, cont)
        if isinstance(b, H.Liste):
            return self.liste(b, cont)
        if isinstance(b, H.Tableau):
            return self.tableau(b, cont)
        if isinstance(b, H.Figure):
            return self.figure(b, cont)
        if isinstance(b, H.Encadre):
            return self.encadre(b, cont)
        if isinstance(b, H.Activite):
            return self.activite(b, cont)
        if isinstance(b, H.Capsule):
            return self.capsule(b)
        if isinstance(b, H.Revision):
            return self.revision(b)
        if isinstance(b, H.Code):
            return self.code(b, cont)
        raise TypeError(f"bloc inconnu : {b!r}")

    def corps(self, cont):
        return cont is self.doc

    # ── intertitres et paragraphes ─────────────────────────────────────────
    def titre(self, b, cont):
        texte = self.typo(H.texte_brut(b.contenu))
        if b.niveau == 2 and self.corps(cont):
            par = self.nouveau_p(cont, "Poly Section")
            m = re.match(r"(\d+)\.\s+(.*)", texte)
            if m:
                r = par.add_run(m.group(1) + ".")
                r.font.color.rgb = RGBColor.from_string(G.ACCENT)
                par.add_run(NBSP + NBSP + m.group(2))
            else:
                par.add_run(texte)
            self.index_section, self.note_section = set(), False
        elif b.niveau <= 3 and self.corps(cont):
            par = self.nouveau_p(cont, "Poly Sous-section")
            par.add_run(texte)
        else:
            par = self.nouveau_p(cont, "Poly Intertitre d'encadré")
            par.add_run(texte)
        if b.ident:
            self.signet(par, b.ident)
        if self.corps(cont):
            self.dernier_p = par
        return par

    def paragraphe(self, b, cont):
        corps = self.corps(cont)
        par = self.nouveau_p(cont, self.style_texte[-1] if not corps else "Normal")
        base = {}
        if b.genre == "chapeau" and corps:
            par.paragraph_format.space_after = Pt(7)
            base = {"taille": 11}
        note, liens = self.ecrire(par, b.contenu, base=base, noter=corps)
        self.poser_marges(par, corps, note, liens)
        if corps:
            self.dernier_p = par
        return par

    def liste(self, b, cont, niveau=0, num=None):
        corps = self.corps(cont)
        if num is None:
            num = self.num.nouvelle_liste_numerotee() if b.ordonnee else self.num.puce
        for contenu, sous in b.items:
            par = self.nouveau_p(cont, "Poly Liste" if corps else self.style_texte[-1])
            G.Numerotation.appliquer(par, num, niveau)
            if not corps:
                par.paragraph_format.space_after = Pt(2)
            note, liens = self.ecrire(par, contenu, noter=corps)
            self.poser_marges(par, corps, note, liens)
            if corps:
                self.dernier_p = par
            if sous:
                self.liste(sous, cont, niveau + 1,
                           num if not sous.ordonnee else self.num.nouvelle_liste_numerotee())

    # ── tableaux ───────────────────────────────────────────────────────────
    def tableau(self, b, cont):
        largeur = self.largeur[-1]
        lignes = b.entetes + b.lignes
        if not lignes:
            return None
        ncol = max(len(l) for l in lignes)
        largeurs = largeurs_colonnes(b, ncol, largeur)
        ancre_tableau = None
        if b.legende:
            par = self.nouveau_p(cont, "Poly Légende")
            par.paragraph_format.space_before = Pt(6)
            par.paragraph_format.space_after = Pt(3)
            par.paragraph_format.keep_with_next = True
            r = par.add_run(self.L["tableau"])
            G._police(r._r.get_or_add_rPr(), G.TITRE)
            r.font.color.rgb = RGBColor.from_string(G.PRIMAIRE)
            par.add_run("  ")
            self.ecrire(par, b.legende, indexer=False, liens=False)
            ancre_tableau = par
        elif self.corps(cont):
            ancre_tableau = self.nouveau_p(cont, "Poly Ancre")
        a_poser = []                       # les QR codes des liens lus dans les cellules : sur l'ancre du tableau
        t = cont.add_table(rows=len(lignes), cols=ncol)
        X.tableau_sans_bordures(t)
        X.tableau_largeurs(t, largeurs)
        fin = len(lignes) - 1
        for i, rang in enumerate(lignes):
            entete = i < len(b.entetes)
            if entete:
                X.ligne_entete(t.rows[i])
            X.ligne_insecable(t.rows[i])
            for j in range(ncol):
                cell = t.rows[i].cells[j]
                X.marges_cellule(cell, 0.07, 0.12, 0.07, 0.12)
                haut = X.filet(sz=8, color=G.PRIMAIRE) if i == 0 else (
                    X.filet(sz=6, color=G.PRIMAIRE) if i == len(b.entetes) and b.entetes
                    else X.filet(sz=4, color=G.FILET))
                bas = X.filet(sz=8, color=G.PRIMAIRE) if i == fin else None
                X.bordures_cellule(cell, top=haut, bottom=bas)
                if entete:
                    X.trame_cellule(cell, G.FOND_PRIMAIRE)
                c = rang[j] if j < len(rang) else H.Cellule([])
                for contenu in (c.contenu or [[]]):
                    par = self.nouveau_p(cell, "Poly Cellule d'en-tête" if entete else "Poly Cellule")
                    base = {"gras": True} if (c.entete and not entete) else {}
                    _, marge = self.ecrire(par, contenu, base=base, indexer=not entete)
                    a_poser.extend(m for m in marge if m[0] not in [x[0] for x in a_poser])
        cible = ancre_tableau if ancre_tableau is not None else self.cible_marge()
        if cible is not None:
            for url, texte, _ in a_poser:
                self.marge(cible, url, self.L["marge"]["en_ligne"], texte, "ecran")
        if self.corps(cont):
            self.nouveau_p(cont, "Poly Espace")
        return t

    # ── figures ────────────────────────────────────────────────────────────
    def image(self, cont, png, l, h, alt, legende, large=False):
        """Une image et sa légende ; large : dans le corps, elle déborde dans la grande marge (17,3 cm)."""
        corps = self.corps(cont)
        largeur = (G.LARGE - 0.02) if (large and corps) else self.largeur[-1]
        hauteur = largeur * h / l
        debord = max(0.0, largeur - G.TEXTE) if (large and corps) else 0.0
        par = self.nouveau_p(cont, "Poly Figure")
        if debord:
            X.retrait_miroir(par)
            par.paragraph_format.right_indent = Cm(-debord)
        par._p.append(X.xml(X.run_image_xml(self.rid(self.imprimable(png, largeur)), largeur, hauteur, self.ids(),
                                            self.L["illustration"], alt, titre="poly:schema|" if debord else "")))
        leg = None
        if legende:
            leg = self.nouveau_p(cont, "Poly Légende")
            if debord:
                X.retrait_miroir(leg)
                leg.paragraph_format.right_indent = Cm(-debord)
            self.ecrire(leg, legende, indexer=False, liens=False)
        if corps:
            self.dernier_p = leg or par
        self.journal.append(dict(genre="figure", largeur=round(largeur, 2), png=png, lieu=self.lieu()))
        return par

    def figure(self, b, cont):
        if not b.source or not os.path.isfile(b.source):
            source = os.path.relpath(b.source, SITE_LOCAL).replace(os.sep, "/") if b.source else "(adresse vide)"
            self.journal.append(dict(genre="image absente", source=source, lieu=self.lieu()))
            return None
        if b.genre == "schema":
            png, l, h = self.atelier.schema_fichier(b.source)
            self.atelier.photographier()
            return self.image(cont, png, l, h, b.alt, b.legende, large=True)
        png, l, h = self.atelier.illustration(b.source)
        return self.image(cont, png, l, h, b.alt, b.legende)

    # ── encadrés ───────────────────────────────────────────────────────────
    def boite(self, genre, cont):
        """Un encadré (classe Boite) ; son paragraphe d'ancrage devient self.ancre, pour les zones de la marge."""
        return Boite(self, genre, cont)

    def etiquette(self, cell, genre, texte, titre=None, picto=None, couleur=None):
        st = STYLES_ENCADRES.get(genre, {})
        couleur = couleur or st.get("couleur", G.PRIMAIRE)
        par = self.nouveau_p(cell, "Poly Étiquette")
        par._p.append(self.run_picto(picto or st.get("picto", "crayon"), couleur, 0.4, -1.6))
        r = par.add_run(NBSP + NBSP + self.typo(texte).upper())
        r.font.color.rgb = RGBColor.from_string(couleur)
        if titre:
            tp = self.nouveau_p(cell, "Poly Intertitre d'encadré")
            if isinstance(titre, str):
                tp.add_run(self.typo(titre))
            else:
                self.ecrire(tp, titre, indexer=False, liens=False)
            return tp
        return par

    def encadre(self, b, cont):
        E = self.L["encadres"]
        genre = b.genre
        bx = self.boite(genre, cont)
        if genre == "retenir":
            lp = self.etiquette(bx.entete, genre, H.texte_brut(b.titre) or E["retenir"])
            if b.etiquette:
                self.signet(lp, b.etiquette)
        elif genre in ("repere", "piege"):
            self.etiquette(bx.entete, genre, E[genre], titre=b.titre or None)
        elif genre == "plus_loin":
            titre = H.texte_brut(b.titre)
            self.etiquette(bx.entete, genre, E["plus_loin"], titre=(titre[:1].upper() + titre[1:]) or None)
        elif genre == "lecture":
            self.etiquette(bx.entete, genre, H.texte_brut(b.titre) or E["lecture"])
        if b.blocs:
            self.style_texte.append("Poly Encadré")
            self.blocs(b.blocs, bx.ligne())
            self.style_texte.pop()
        bx.fermer()
        return bx.table

    def code(self, b, cont):
        if re.search(r"https?://|\bwww\.", b.texte):
            return self.code_en_ligne(b, cont)
        par = self.nouveau_p(cont, "Poly Code")
        par.paragraph_format.space_before = Pt(2)
        par.paragraph_format.space_after = Pt(8)
        X.trame_paragraphe(par, G.FOND_GRIS)
        X.bordures_paragraphe(par, left=X.filet(sz=12, space=6, color=G.PRIMAIRE))
        par.paragraph_format.left_indent = Cm(0.25)
        for k, l in enumerate(b.texte.split("\n")):
            if k:
                par.add_run().add_break()
            par.add_run(l.replace("\t", "    "))
        if self.corps(cont):
            self.dernier_p = par
        return par

    def code_en_ligne(self, b, cont):
        """Un code qui porte une adresse ne s'imprime pas : aucune adresse dans le texte. Une phrase dit où le copier,
        et le QR code de la marge ouvre la page."""
        texte = ("Ce code contient une adresse : copiez-le depuis le site, avec le QR code de la marge."
                 if self.langue.code == "fr" else
                 "This code contains a web address: copy it from the site, using the QR code in the margin.")
        par = self.nouveau_p(cont, "Poly Code")
        par.paragraph_format.space_before = Pt(2)
        par.paragraph_format.space_after = Pt(8)
        X.trame_paragraphe(par, G.FOND_GRIS)
        X.bordures_paragraphe(par, left=X.filet(sz=12, space=6, color=G.PRIMAIRE))
        par.paragraph_format.left_indent = Cm(0.25)
        r = par.add_run(texte)
        G._police(r._r.get_or_add_rPr(), G.SANS)
        r.font.size = Pt(8.5)
        r.font.italic = True
        url = self.langue.url(self.page.rel) if self.page else self.langue.url("")
        cible = par if self.corps(cont) else self.cible_marge()
        if cible is not None:
            self.marge(cible, url, self.L["marge"]["en_ligne"], "", "code")
        if self.corps(cont):
            self.dernier_p = par
        return par

    # ── activités et capsules ──────────────────────────────────────────────
    def activite(self, b, cont):
        d, P, E, M = b.donnees, self.L["papier"], self.L["encadres"], self.L["marge"]
        page_url = self.langue.url(self.page.rel) if self.page else self.langue.url("")
        if b.genre == "checkpoint":
            bx = self.boite("controle", cont)
            self.signer(b.ident)
            titre = d.get("titre", "")
            libelle, _, sujet = titre.partition(" — ")
            self.etiquette(bx.entete, "controle", libelle or E["controle"],
                           titre=(sujet[:1].upper() + sujet[1:]) or None)
            self.consigne(bx.entete, P["controle"])
            for k, q in enumerate(d.get("questions", []), 1):
                cell = bx.ligne(insecable=True)
                pq = self.nouveau_p(cell, "Poly Encadré")
                pq.paragraph_format.space_before = Pt(4)
                pq.paragraph_format.space_after = Pt(2)
                self.numero(pq, k)
                self.ecrire(pq, markdown(q.get("q", "")), indexer=False, liens=False)
                self.propositions(cell, q.get("options", []))
            self.marge(self.ancre, page_url + "#" + b.ident, M["corrige"], court(titre) or E["controle"], "coche")
            bx.fermer()
            return bx.table
        if b.genre == "reveal":
            bx = self.boite("situation", cont)
            self.signer(b.ident)
            self.etiquette(bx.entete, "situation", E["situation"])
            options = d.get("options") or []
            self.consigne(bx.entete, P["reveal_choix"] if options else P["reveal"])
            cell = bx.ligne(insecable=True)
            pq = self.nouveau_p(cell, "Poly Encadré")
            self.ecrire(pq, markdown(d.get("question", "")), indexer=False, liens=False)
            if options:
                self.propositions(cell, options)
            else:
                for _ in range(3):
                    pl = self.nouveau_p(cell, "Poly Ligne d'écriture")
                    pointille = X.filet(val="dotted", sz=6, space=1, color=G.ARDOISE)
                    X.bordures_paragraphe(pl, bottom=pointille, between=pointille)
            if d.get("indice"):
                self.indice(cell, d["indice"])
            self.marge(self.ancre, page_url + "#" + b.ident, M["reponse"], E["situation"].capitalize(), "question")
            bx.fermer()
            return bx.table
        if b.genre == "match":
            bx = self.boite("activite", cont)
            self.signer(b.ident)
            self.etiquette(bx.entete, "activite", E["activite"], titre=d.get("titre"))
            self.consigne(bx.entete, papier(d.get("consigne", ""), P, P["match"]))
            cell = bx.ligne()
            paires = d.get("paires", [])
            droite = [p.get("d", "") for p in paires]
            random.Random(b.ident).shuffle(droite)          # mélangée : l'ordre ne souffle rien
            largeur = self.largeur[-1]
            tg = self.tableau_cellule(cell, 1 + len(paires), 2)
            X.tableau_sans_bordures(tg)
            X.tableau_largeurs(tg, [largeur / 2, largeur / 2])
            for j, titre in enumerate((d.get("colonneGauche", ""), d.get("colonneDroite", ""))):
                c = tg.rows[0].cells[j]
                X.bordures_cellule(c, bottom=X.filet(sz=6, color=G.PRIMAIRE))
                self.nouveau_p(c, "Poly Cellule d'en-tête").add_run(self.typo(titre))
            for i, p in enumerate(paires, 1):
                for j, texte in enumerate((f"{i}.  " + p.get("g", ""), "ABCDEFGHIJ"[i - 1] + ".  " + droite[i - 1])):
                    c = tg.rows[i].cells[j]
                    X.bordures_cellule(c, bottom=X.filet(sz=4, color=G.FILET))
                    X.marges_cellule(c, 0.05, 0.08, 0.05, 0.12)
                    self.ecrire(self.nouveau_p(c, "Poly Cellule"), markdown(texte), indexer=False, liens=False)
            self.marge(self.ancre, page_url + "#" + b.ident, M["activite"], court(d.get("titre", "")), "crayon")
            bx.fermer()
            return bx.table
        if b.genre == "spotError":
            bx = self.boite("activite", cont)
            self.signer(b.ident)
            self.etiquette(bx.entete, "activite", E["erreurs"], titre=d.get("titre"))
            self.consigne(bx.entete, papier(d.get("consigne", ""), P, P["erreurs"]))
            if d.get("intro"):
                self.ecrire(self.nouveau_p(bx.entete, "Poly Encadré"), markdown(d["intro"]), indexer=False,
                            liens=False)
            cell = bx.ligne()
            for k, f in enumerate(d.get("fragments", []), 1):
                po = self.nouveau_p(cell, "Poly Encadré")
                po.paragraph_format.left_indent = Cm(0.55)
                po.paragraph_format.first_line_indent = Cm(-0.55)
                po.paragraph_format.space_after = Pt(2)
                po.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self.case(po)
                rl = po.add_run(NBSP + str(k) + "." + NBSP + NBSP)
                G._police(rl._r.get_or_add_rPr(), G.TITRE)
                rl.font.size = Pt(8.5)
                rl.font.color.rgb = RGBColor.from_string(G.PRIMAIRE)
                self.ecrire(po, markdown(f.get("t", "")), indexer=False, liens=False)
            self.marge(self.ancre, page_url + "#" + b.ident, M["activite"], E["erreurs"].capitalize(), "crayon")
            bx.fermer()
            return bx.table
        if b.genre == "branch":
            bx = self.boite("activite", cont)
            self.signer(b.ident)
            self.etiquette(bx.entete, "activite", E["activite"], titre=d.get("titre"))
            self.consigne(bx.entete, P["branch"])
            if d.get("intro"):
                self.ecrire(self.nouveau_p(bx.entete, "Poly Encadré"), markdown(d["intro"]), indexer=False,
                            liens=False)
            for k, etape in enumerate(d.get("etapes", []), 1):
                cell = bx.ligne(insecable=True)
                ps = self.nouveau_p(cell, "Poly Encadré")
                ps.paragraph_format.space_before = Pt(4)
                if len(d.get("etapes", [])) > 1:
                    self.numero(ps, k)
                self.ecrire(ps, markdown((etape.get("situation", "") + " " + etape.get("question", "")).strip()),
                            indexer=False, liens=False)
                self.propositions(cell, [o.get("t", "") for o in (etape.get("options") or etape.get("choix") or [])])
            self.marge(self.ancre, page_url + "#" + b.ident, M["activite"], court(d.get("titre", "")), "crayon")
            bx.fermer()
            return bx.table
        if b.genre == "hotspots":
            return self.schema_zones(b, cont, page_url)
        return None

    def schema_zones(self, b, cont, page_url):
        """Un schéma à zones cliquables : l'image, ses zones devenues des pastilles numérotées, sa légende, puis
        « Lire ce schéma, élément par élément » ; le schéma interactif en QR code dans la marge."""
        d, E, M = b.donnees, self.L["encadres"], self.L["marge"]
        src = (d.get("src") or "").strip()               # vide : un schéma pas encore produit (squelette)
        chemin = self.langue.locale(self.langue.resoudre(src, self.page.rel)) if (self.page and src) else None
        points = d.get("points") or []
        if not chemin or not os.path.isfile(chemin):
            self.journal.append(dict(genre="schema absent", source=src or "(adresse vide)", lieu=self.lieu()))
            return None
        # Un schéma qui numérote déjà ses zones (« 1. Le public ») garde son image telle quelle : une pastille
        # cacherait son texte, et l'encadré de lecture reprend ses numéros.
        numerotes = bool(points) and all(re.match(r"\s*\d+\s*[.)]", p.get("titre", "")) for p in points)
        png, l, h = self.atelier.schema_zones(chemin, [] if numerotes else points)
        par = self.image(cont, png, l, h, d.get("alt", ""), markdown(d.get("legende", "")), large=True)
        self.signet(par, b.ident)
        cible = self.dernier_p if self.corps(cont) else self.cible_marge()
        if cible is not None:
            self.marge(cible, page_url + "#" + b.ident, M["en_ligne"], M["schema"], "oeil")
        if points:
            bx = self.boite("lecture", cont)
            self.etiquette(bx.entete, "lecture", E["lecture"])
            cell = bx.ligne()
            for k, p in enumerate(points, 1):
                pp = self.nouveau_p(cell, "Poly Encadré")
                pp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pp.paragraph_format.space_after = Pt(3)
                if not numerotes:
                    self.numero(pp, k)
                r = pp.add_run(self.typo(p.get("titre", "")))
                r.font.bold = True
                if p.get("texte"):
                    pp.add_run(NBSP + "—" + " ")
                    self.ecrire(pp, markdown(p["texte"]), indexer=False, liens=False)
            bx.fermer()
        return par

    def numero(self, par, k):
        r = par.add_run(f"{k}.")
        G._police(r._r.get_or_add_rPr(), G.TITRE)
        r.font.color.rgb = RGBColor.from_string(G.PRIMAIRE)
        par.add_run(NBSP + NBSP)

    def propositions(self, cell, options):
        """Les propositions d'une question, chacune précédée d'une case vide et de sa lettre."""
        for i, o in enumerate(options):
            po = self.nouveau_p(cell, "Poly Encadré")
            po.paragraph_format.left_indent = Cm(0.55)
            po.paragraph_format.space_after = Pt(1)
            po.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self.case(po)
            rl = po.add_run(NBSP + "ABCDEFGHIJ"[i] + NBSP + NBSP)
            G._police(rl._r.get_or_add_rPr(), G.TITRE)
            rl.font.size = Pt(8.5)
            rl.font.color.rgb = RGBColor.from_string(G.PRIMAIRE)
            self.ecrire(po, markdown(o), indexer=False, liens=False)

    def consigne(self, cell, texte):
        par = self.nouveau_p(cell, "Poly Petit texte")
        par.paragraph_format.space_after = Pt(4)
        r = par.add_run(self.typo(texte))
        r.font.italic = True
        r.font.color.rgb = RGBColor.from_string(G.ARDOISE)
        return par

    def indice(self, cell, texte):
        par = self.nouveau_p(cell, "Poly Petit texte")
        par.paragraph_format.space_before = Pt(5)
        par.paragraph_format.space_after = Pt(0)
        r = par.add_run(self.L["papier"]["indice"] + (NBSP + ":" + NBSP if self.langue.code == "fr" else ": "))
        G._police(r._r.get_or_add_rPr(), G.TITRE)
        r.font.color.rgb = RGBColor.from_string(G.PRIMAIRE)
        self.ecrire(par, markdown(texte), base={"italique": True}, indexer=False, liens=False)

    def case(self, par, taille=10):
        r = par.add_run(CASE)
        G._police(r._r.get_or_add_rPr(), SYMBOLE)
        r.font.size = Pt(taille)
        r.font.color.rgb = RGBColor.from_string(G.PRIMAIRE)
        return r

    def capsule(self, b, par=None):
        d, M = b.donnees, self.L["marge"]
        cible = par or self.cible_marge()
        if cible is None or not d.get("src"):
            return None
        url = self.langue.resoudre(d.get("src", ""), self.page.rel) if self.page else d["src"]
        etiquette = M["ecouter"] + (" · " + duree(d.get("duree", ""), M) if d.get("duree") else "")
        return self.marge(cible, url, etiquette, d.get("titre", ""), "casque")

    def revision(self, b):
        """Les flashcards du chapitre : un QR code dans la marge, vers le paquet de la page."""
        cible = self.cible_marge()
        if cible is None or self.page is None:
            return None
        return self.marge(cible, self.langue.url(self.page.rel) + "#" + b.ident, self.L["marge"]["reviser"],
                          self.L["marge"]["flashcards"], "cartes")


# ── l'encadré ──────────────────────────────────────────────────────────────
def reecrire(contenu, regles):
    """Applique des substitutions (motif, remplacement) à un texte en ligne, par-dessus ses morceaux : le texte
    remplacé prend la forme du premier caractère qu'il remplace, sans lien."""
    plats = [(c, s) for s in contenu for c in s.texte]
    texte = "".join(c for c, _ in plats)
    for motif, rempl in regles:
        if not re.search(motif, texte):
            continue
        out, pos = [], 0
        for m in re.finditer(motif, texte):
            out.extend(plats[pos:m.start()])
            modele = plats[m.start()][1] if m.start() < len(plats) else H.Seg("")
            neuf = H.Seg("", modele.gras, modele.italique, modele.code, None, modele.exposant)
            out.extend((c, neuf) for c in rempl)
            pos = m.end()
        out.extend(plats[pos:])
        plats = out
        texte = "".join(c for c, _ in plats)
    segs = []
    for c, s in plats:
        if segs and segs[-1][1] is s:
            segs[-1][0].append(c)
        else:
            segs.append(([c], s))
    return [H.Seg("".join(cs), s.gras, s.italique, s.code, s.lien, s.exposant) for cs, s in segs]


class Boite:
    """Un encadré : un tableau d'une colonne. Sa première ligne, l'en-tête (pictogramme, libellé, titre, consigne),
    reste avec la suivante : ses paragraphes sont « liés au suivant », et Word lie alors la ligne entière à la ligne
    qui suit. Les lignes de contenu n'ont aucun paragraphe lié : elles se coupent d'une page à l'autre (une
    question de point de contrôle, insécable, tient sur sa page). Filets et fond se posent à la fermeture : le
    filet du haut sur la première ligne, celui du bas sur la dernière, les filets latéraux et le fond partout."""

    def __init__(self, rendu, genre, cont):
        self.r, self.genre, self.cont = rendu, genre, cont
        self.st = STYLES_ENCADRES[genre]
        self.dans_corps = rendu.corps(cont)
        if self.dans_corps:
            rendu.ancre = rendu.nouveau_p(cont, "Poly Ancre")
        self.largeur = rendu.largeur[-1]
        self.table = cont.add_table(rows=1, cols=1)
        X.tableau_sans_bordures(self.table)
        X.tableau_largeurs(self.table, [self.largeur])
        X.ligne_insecable(self.table.rows[0])
        h, g, b, d = self.st["marges"]
        rendu.largeur.append(self.largeur - g - d)
        self.entete = self.table.rows[0].cells[0]
        rendu.cellules_entete.add(self.entete._tc)

    def ligne(self, insecable=False):
        """Une nouvelle ligne de contenu ; rend sa cellule."""
        rang = self.table.add_row()
        if insecable:
            X.ligne_insecable(rang)
        cell = rang.cells[0]
        X.largeur_cellule(cell, self.largeur)
        return cell

    def fermer(self, espace=True):
        lignes = self.table.rows
        n = len(lignes)
        h, g, b, d = self.st["marges"]
        for i, rang in enumerate(lignes):
            cell = rang.cells[0]
            self.r.fin_de_cellule(cell)
            haut = h if i == 0 else 0.02
            bas = b if i == n - 1 else 0.02
            X.marges_cellule(cell, haut, g, bas, d)
            bords = dict(self.st["bords"])
            if i > 0:
                bords.pop("top", None)
            if i < n - 1:
                bords.pop("bottom", None)
            X.bordures_cellule(cell, **bords)
            if self.st["fond"]:
                X.trame_cellule(cell, self.st["fond"])
        self.r.largeur.pop()
        if self.dans_corps:
            if espace:
                self.r.nouveau_p(self.cont, "Poly Espace")
            self.r.ancre = None


# ── utilitaires ────────────────────────────────────────────────────────────
def court(titre, limite=34):
    """Une légende de QR code courte : le titre, ou sa partie avant le tiret, ou rien (l'encadré voisin le porte)."""
    t = (titre or "").strip()
    if len(t) <= limite:
        return t
    tete = t.split(" — ")[0].strip()
    return tete if len(tete) <= limite else ""


def signet_nom(ident):
    return re.sub(r"[^A-Za-z0-9_]", "_", "poly_" + ident)[:40]


def capitale(legende):
    """Une légende de QR code commence par une capitale — sauf un nom qui s'écrit en minuscule (une adresse)."""
    premier = legende.split(" ")[0] if legende else ""
    if not premier or "." in premier:
        return legende
    return legende[:1].upper() + legende[1:]


def vedette(terme, L):
    """Un terme du glossaire tel que le livret l'imprime en vedette (glossaire, index, définition dans la marge) : une
    capitale initiale, dans les deux langues — sauf aux noms qui s'écrivent en minuscule (« minuscules » de la fiche
    du livret)."""
    if not terme or not terme[:1].islower() or terme.split()[0] in L.get("minuscules", ()):
        return terme
    return terme[:1].upper() + terme[1:]


def duree(texte, M):
    m = re.match(r"(\d+):(\d+)", texte or "")
    return M["duree"].format(m=int(m.group(1)), s=int(m.group(2))) if m else texte


def markdown(texte):
    """Le gras **…** des données des activités → texte en ligne."""
    out = []
    for k, morceau in enumerate(re.split(r"\*\*", texte or "")):
        if morceau:
            out.append(H.Seg(morceau, gras=bool(k % 2)))
    return H.replier(out)


def papier(consigne, P, tete):
    """La consigne d'une activité, pour le papier : la phrase qui parle de cliquer devient celle du crayon."""
    phrases = re.split(r"(?<=[.!?])\s+", (consigne or "").strip())
    gardees = [p for p in phrases if p and not any(m in p.lower() for m in P["clic"])]
    return " ".join([tete] + gardees)


def largeurs_colonnes(b, ncol, largeur):
    """Des largeurs de colonnes à la mesure de leur texte : plus de place au texte long, jamais moins que le plus
    long mot (0,19 cm par caractère, gras compris, et les marges de la cellule)."""
    longueurs, mots = [0.0] * ncol, [0] * ncol
    for rang in b.entetes + b.lignes:
        for j, c in enumerate(rang[:ncol]):
            t = " ".join(H.texte_brut(p) for p in c.contenu)
            longueurs[j] = max(longueurs[j], 0.0) + len(t)
            mots[j] = max([mots[j]] + [len(m) for m in t.split()])
    n = max(1, len(b.entetes + b.lignes))
    poids = [max(8.0, l / n) ** 0.85 for l in longueurs]
    total = sum(poids)
    larg = [largeur * p / total for p in poids]
    minimum = [min(largeur / ncol, 0.19 * m + 0.42) for m in mots]
    for _ in range(3):
        manque = sum(max(0, mi - l) for mi, l in zip(minimum, larg))
        if manque <= 0:
            break
        surplus = [max(0, l - mi) for mi, l in zip(minimum, larg)]
        s = sum(surplus) or 1
        larg = [max(mi, l - manque * su / s) for mi, l, su in zip(minimum, larg, surplus)]
    k = largeur / sum(larg)
    return [l * k for l in larg]

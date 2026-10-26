# -*- coding: utf-8 -*-
"""La mise en page du polycopié : format, marges en miroir, styles, puces, sections, en-têtes et pieds.

LA PAGE. A4 portrait, marges en miroir : intérieure 2,2 cm, extérieure 5,2 cm (la marge des QR codes et des notes),
haut 2,2 cm, bas 2,4 cm. Le texte courant tient dans une colonne de 13,6 cm ; un schéma déborde de 3,7 cm dans la
grande marge (LARGE, 17,3 cm), jusqu'au même bord que le filet de l'en-tête. Les pages de référence (glossaire,
index) prennent toute la largeur : marges de 2,2 cm des deux côtés.

LES CARACTÈRES (installés sur le poste) : le texte en Cambria 10,5 pt, interligne 1,2, justifié et coupé ; les
titres en Segoe UI Semibold ; la marge et les légendes en Segoe UI 7 à 8 pt ; le code en Consolas 9 pt.

LES COULEURS, celles du site : le bleu nuit (primaire) et le doré (accent), des fonds très clairs. Rien ne se lit à
la seule couleur : un encadré se reconnaît à son pictogramme, à son libellé et à son filet, même imprimé en noir et
blanc. Le doré (ACCENT) ne sert qu'aux numéros et aux libellés en gros caractères ; le petit texte doré prend
ACCENT_FONCE (contraste de 6 sur blanc).

LES EN-TÊTES : le titre du livre sur les pages paires, « Chapitre n · titre » sur les pages impaires (deux champs
STYLEREF, sur les styles « Poly Numéro de chapitre » et « Poly Titre de chapitre »), filet fin ; le numéro de page au
coin extérieur. Les pages d'ouverture n'ont ni en-tête ni pied.
"""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import ooxml as X

# ── les mesures (cm) ───────────────────────────────────────────────────────
PAGE_L, PAGE_H = 21.0, 29.7
INTERIEURE, EXTERIEURE, HAUT, BAS = 2.2, 5.2, 2.2, 2.4
EN_TETE, PIED = 1.1, 1.2                      # distance du bord à l'en-tête et au pied
TEXTE = PAGE_L - INTERIEURE - EXTERIEURE      # 13,6 cm : la colonne de texte
DEBORD = 3.7                                  # ce qu'un schéma prend sur la grande marge
LARGE = TEXTE + DEBORD                        # 17,3 cm
REFERENCE_MARGE = 2.2                         # pages de référence : marges symétriques
REFERENCE = PAGE_L - 2 * REFERENCE_MARGE      # 16,6 cm
HAUTEUR_UTILE = PAGE_H - HAUT - BAS           # 25,1 cm
MARGE_BOITE = 3.6                             # largeur d'une zone de la marge (QR code et légende)
QR = 2.2                                      # côté d'un QR code de la marge

# ── les couleurs ───────────────────────────────────────────────────────────
PRIMAIRE = "1B3A5B"
PRIMAIRE_FONCE = "0E2238"
ACCENT = "9C7A22"
ACCENT_FONCE = "7A5E17"
ENCRE = "16202E"
ARDOISE = "4A5A6E"
FILET = "C9D3DE"
FOND_PRIMAIRE = "EAF0F6"
FOND_ACCENT = "FAF4E6"
FOND_GRIS = "F3F5F8"
FOND_PIEGE = "FAEDEA"
BRIQUE = "A93E2C"
BLANC = "FFFFFF"

# ── les polices ────────────────────────────────────────────────────────────
TEXTE_POLICE = "Cambria"
TITRE = "Segoe UI Semibold"
SANS = "Segoe UI"
LEGER = "Segoe UI Light"
CODE = "Consolas"


def _police(style_ou_run, police):
    """La police sur les quatre jeux de caractères (sinon Word garde Calibri pour certains signes)."""
    rpr = style_ou_run.element.get_or_add_rPr() if hasattr(style_ou_run, "element") else style_ou_run
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = X.xml("<w:rFonts/>")
        X.poser(rpr, rf)
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn("w:" + a), police)
    for a in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        rf.attrib.pop(qn("w:" + a), None)


# ── les styles ─────────────────────────────────────────────────────────────
# nom affiché, identifiant, police, taille (pt), couleur, gras, italique, avant, après (pt), interligne (multiple ou
# ('exact', pt)), alignement, niveau de plan, garder avec le suivant, options.
STYLES = [
    ("Poly Numéro de chapitre", "PolyNumeroChapitre", TITRE, 11, ACCENT, False, False, 0, 0, 1.0, "left", None, True,
     {}),
    ("Poly Titre de chapitre", "PolyTitreChapitre", TITRE, 26, ENCRE, False, False, 2, 8, ("exact", 31), "left", 0,
     True, {"coupure": False}),
    ("Poly Titre de partie", "PolyTitrePartie", TITRE, 24, PRIMAIRE, False, False, 0, 6, ("exact", 30), "left", 0,
     True, {"coupure": False}),
    ("Poly Titre liminaire", "PolyTitreLiminaire", TITRE, 19, PRIMAIRE, False, False, 0, 6, ("exact", 24), "left", 0,
     True, {"coupure": False}),
    ("Poly Titre sans plan", "PolyTitreSansPlan", TITRE, 19, PRIMAIRE, False, False, 0, 6, ("exact", 24), "left",
     None, True, {"coupure": False}),
    ("Poly Section", "PolySection", TITRE, 14, PRIMAIRE, False, False, 18, 6, ("exact", 18), "left", 1, True,
     {"coupure": False, "lignes": True}),
    ("Poly Sous-section", "PolySousSection", TITRE, 11.5, ENCRE, False, False, 12, 4, ("exact", 15), "left", 2, True,
     {"coupure": False, "lignes": True}),
    ("Poly Intertitre d'encadré", "PolyIntertitre", TITRE, 10, ENCRE, False, False, 0, 3, ("exact", 13), "left",
     None, True, {"coupure": False}),
    ("Poly Étiquette", "PolyEtiquette", TITRE, 7.5, PRIMAIRE, False, False, 0, 3, ("min", 11), "left", None, True,
     {"coupure": False, "interlettre": 0.8}),
    ("Poly Encadré", "PolyEncadre", TEXTE_POLICE, 10, ENCRE, False, False, 0, 4, 1.15, "both", None, False, {}),
    ("Poly Liste", "PolyListe", TEXTE_POLICE, 10.5, ENCRE, False, False, 0, 3, 1.2, "both", None, False, {}),
    ("Poly Légende", "PolyLegende", SANS, 8, ARDOISE, False, False, 3, 10, ("exact", 10.5), "left", None, False,
     {"coupure": False}),
    ("Poly Figure", "PolyFigure", TEXTE_POLICE, 10, ENCRE, False, False, 6, 0, 1.0, "center", None, True, {}),
    ("Poly Code", "PolyCode", CODE, 9, ENCRE, False, False, 0, 0, ("exact", 11.5), "left", None, False,
     {"coupure": False}),
    ("Poly Cellule", "PolyCellule", TEXTE_POLICE, 9, ENCRE, False, False, 0, 2, 1.1, "left", None, False,
     {"coupure": False}),
    ("Poly Cellule d'en-tête", "PolyCelluleEntete", TITRE, 8.5, PRIMAIRE, False, False, 0, 0, 1.05, "left", None,
     False, {"coupure": False}),
    ("Poly Source", "PolySource", SANS, 7.5, ARDOISE, False, False, 4, 0, ("exact", 10), "left", None, False,
     {"coupure": False}),
    ("Poly Marge", "PolyMarge", SANS, 7, ARDOISE, False, False, 0, 0, ("exact", 8.5), "center", None, False,
     {"coupure": False}),
    ("Poly En-tête", "PolyEnTete", SANS, 7.5, ARDOISE, False, False, 0, 0, 1.0, "left", None, False,
     {"coupure": False}),
    ("Poly Pied", "PolyPied", TITRE, 9, PRIMAIRE, False, False, 0, 0, 1.0, "left", None, False, {"coupure": False}),
    ("Poly Mini-sommaire", "PolyMiniSommaire", SANS, 9.5, ENCRE, False, False, 0, 3, ("exact", 13.5), "left", None,
     False, {"coupure": False}),
    ("Poly Petit texte", "PolyPetitTexte", SANS, 8.5, ENCRE, False, False, 0, 3, ("exact", 12), "left", None, False,
     {"coupure": False}),
    ("Poly Lettre", "PolyLettre", TITRE, 12, PRIMAIRE, False, False, 10, 4, ("exact", 15), "left", None, True,
     {"coupure": False}),
    # L'ancre d'un encadré (les zones de la marge s'y accrochent) et l'espace qui suit un tableau.
    ("Poly Ancre", "PolyAncre", SANS, 1, ENCRE, False, False, 5, 0, ("exact", 1), "left", None, True,
     {"coupure": False}),
    ("Poly Espace", "PolyEspace", SANS, 1, ENCRE, False, False, 0, 7, ("exact", 1), "left", None, False,
     {"coupure": False}),
    ("Poly Ligne d'écriture", "PolyLigneEcriture", SANS, 9, ARDOISE, False, False, 0, 0, ("exact", 20), "left",
     None, False, {"coupure": False}),
    ("Poly Glossaire terme", "PolyGlossaireTerme", TITRE, 9.5, ENCRE, False, False, 5, 1, ("exact", 12.5), "left",
     None, True, {"coupure": False}),
    ("Poly Glossaire définition", "PolyGlossaireDefinition", TEXTE_POLICE, 9, ENCRE, False, False, 0, 4, 1.12,
     "left", None, False, {}),
    # Les styles que Word reconnaît par leur nom : la table des matières et l'index.
    ("toc 1", "TOC1", TITRE, 10, PRIMAIRE, False, False, 8, 2, 1.0, "left", None, True,
     {"coupure": False, "tab": True}),
    ("toc 2", "TOC2", TEXTE_POLICE, 10, ENCRE, False, False, 0, 2, 1.0, "left", None, False,
     {"coupure": False, "tab": True, "retrait": 0.55}),
    ("index 1", "Index1", TEXTE_POLICE, 9, ENCRE, False, False, 0, 1, 1.0, "left", None, False,
     {"coupure": False, "suspendu": 0.4}),
    ("index heading", "IndexHeading", TITRE, 10, PRIMAIRE, False, False, 8, 2, 1.0, "left", None, True,
     {"coupure": False}),
]


def _interligne(pf, valeur):
    """Un multiple, ('exact', pt) ou ('min', pt) : au moins, pour une ligne qui porte une image."""
    if isinstance(valeur, tuple):
        pf.line_spacing = Pt(valeur[1])
        pf.line_spacing_rule = 4 if valeur[0] == "exact" else 3      # WD_LINE_SPACING.EXACTLY, AT_LEAST
    else:
        pf.line_spacing = valeur


def creer_styles(doc, largeur_tab=TEXTE):
    styles = doc.styles
    normal = styles["Normal"]
    _police(normal, TEXTE_POLICE)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(ENCRE)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(5)
    pf.line_spacing = 1.2
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True
    existants = {s.name: s for s in styles}
    for (nom, ident, police, taille, couleur, gras, italique, avant, apres, interligne, jc, niveau, garder,
         opt) in STYLES:
        integre = not nom.startswith("Poly")           # toc 1, index 1… : les styles que Word connaît par leur nom
        st = existants.get(nom) or styles.add_style(nom, WD_STYLE_TYPE.PARAGRAPH, builtin=integre)
        st.style_id = ident
        st.base_style = normal
        st.quick_style = nom.startswith("Poly")
        _police(st, police)
        st.font.size = Pt(taille)
        st.font.color.rgb = RGBColor.from_string(couleur)
        st.font.bold = gras or None
        st.font.italic = italique or None
        p = st.paragraph_format
        p.space_before, p.space_after = Pt(avant), Pt(apres)
        _interligne(p, interligne)
        p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "both": WD_ALIGN_PARAGRAPH.JUSTIFY, "right": WD_ALIGN_PARAGRAPH.RIGHT}[jc]
        p.keep_with_next = garder
        if opt.get("lignes"):
            p.keep_together = True
        if opt.get("retrait"):
            p.left_indent = Cm(opt["retrait"])
        if opt.get("suspendu"):
            p.left_indent = Cm(opt["suspendu"])
            p.first_line_indent = Cm(-opt["suspendu"])
        pp = st.element.get_or_add_pPr()
        if niveau is not None:
            X.poser(pp, X.xml(f'<w:outlineLvl w:val="{niveau}"/>'))
        if opt.get("coupure") is False:
            X.poser(pp, X.xml("<w:suppressAutoHyphens/>"))
        if opt.get("tab"):
            tabs = X.xml("<w:tabs/>")
            tabs.append(X.xml(f'<w:tab w:val="right" w:leader="dot" w:pos="{X.twips(largeur_tab)}"/>'))
            X.poser(pp, tabs)
        if opt.get("interlettre"):
            X.poser(st.element.get_or_add_rPr(), X.xml(f'<w:spacing w:val="{int(opt["interlettre"] * 20)}"/>'))
    # Les liens hypertextes de la table des matières ne se soulignent pas en bleu.
    if "Hyperlink" in existants:
        h = existants["Hyperlink"]
        h.font.color.rgb = RGBColor.from_string(ENCRE)
        h.font.underline = False


# ── les puces et les numéros ───────────────────────────────────────────────
class Numerotation:
    """Les listes du livret : une puce « • » (niveau 2 : « – ») et des listes numérotées « 1. », chacune repartant de
    1."""

    PUCE = 9001
    NUMERO = 9002

    def __init__(self, doc):
        self.doc = doc
        self.racine = doc.part.numbering_part.element
        self._abstrait(self.PUCE, [("bullet", "•", PRIMAIRE, TEXTE_POLICE), ("bullet", "–", PRIMAIRE, TEXTE_POLICE),
                                   ("bullet", "·", PRIMAIRE, TEXTE_POLICE)])
        self._abstrait(self.NUMERO, [("decimal", "%1.", PRIMAIRE, TITRE), ("lowerLetter", "%2)", PRIMAIRE, TITRE),
                                     ("lowerRoman", "%3.", PRIMAIRE, TITRE)])
        self.puce = self._num(self.PUCE)
        self._suivant = 9100

    def _abstrait(self, ident, niveaux):
        lv = []
        for i, (fmt, texte, couleur, police) in enumerate(niveaux):
            gauche = 0.5 + 0.5 * i
            lv.append(f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
                      f'<w:lvlText w:val="{texte}"/><w:lvlJc w:val="left"/>'
                      f'<w:pPr><w:ind w:left="{X.twips(gauche)}" w:hanging="{X.twips(0.4)}"/></w:pPr>'
                      f'<w:rPr><w:rFonts w:ascii="{police}" w:hAnsi="{police}" w:cs="{police}"/>'
                      f'<w:color w:val="{couleur}"/></w:rPr></w:lvl>')
        el = X.xml(f'<w:abstractNum w:abstractNumId="{ident}"><w:multiLevelType w:val="hybridMultilevel"/>'
                   + "".join(lv) + "</w:abstractNum>")
        premiers = self.racine.findall(qn("w:num"))
        if premiers:
            premiers[0].addprevious(el)
        else:
            self.racine.append(el)

    def _num(self, abstrait, redemarrer=False):
        ident = abstrait if not redemarrer else self._suivant
        corps = f'<w:abstractNumId w:val="{abstrait}"/>'
        if redemarrer:
            corps += '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>'
        self.racine.append(X.xml(f'<w:num w:numId="{ident}">{corps}</w:num>'))
        return ident

    def nouvelle_liste_numerotee(self):
        self._suivant += 1
        return self._num(self.NUMERO, redemarrer=True)

    @staticmethod
    def appliquer(paragraphe, num_id, niveau=0):
        X.poser(X.ppr(paragraphe), X.xml(f'<w:numPr><w:ilvl w:val="{niveau}"/><w:numId w:val="{num_id}"/></w:numPr>'))


# ── le document ────────────────────────────────────────────────────────────
def nouveau_document(langue, titre, sujet):
    """Un document vide aux réglages du livret : A4, marges en miroir, en-têtes pairs et impairs, mode Word 2013 et
    suivants, coupure automatique des mots, langue du texte, propriétés."""
    L = langue.L
    doc = Document()
    reglages = doc.settings.element
    for nom in ("mirrorMargins", "evenAndOddHeaders", "autoHyphenation"):
        X.poser(reglages, X.xml(f"<w:{nom}/>"))
    X.poser(reglages, X.xml('<w:consecutiveHyphenLimit w:val="2"/>'))
    X.poser(reglages, X.xml(f'<w:hyphenationZone w:val="{X.twips(0.6)}"/>'))
    X.poser(reglages, X.xml('<w:doNotHyphenateCaps/>'))
    for c in reglages.iter(qn("w:compatSetting")):
        if c.get(qn("w:name")) == "compatibilityMode":
            c.set(qn("w:val"), "15")
    tfl = reglages.find(qn("w:themeFontLang"))
    if tfl is not None:
        tfl.set(qn("w:val"), L["langue_word"])
    for ls in reglages.findall(qn("w:listSeparator")):
        ls.set(qn("w:val"), ";" if langue.code == "fr" else ",")
    for ds in reglages.findall(qn("w:decimalSymbol")):
        ds.set(qn("w:val"), "," if langue.code == "fr" else ".")
    # La langue du texte, pour la coupure des mots et le tri de l'index.
    defaut = doc.styles.element.find(qn("w:docDefaults"))
    rpr = defaut.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    X.poser(rpr, X.xml(f'<w:lang w:val="{L["langue_word"]}" w:eastAsia="en-US" w:bidi="ar-SA"/>'))
    creer_styles(doc)
    cp = doc.core_properties
    cp.title, cp.subject, cp.author = titre, sujet, L["auteur"]
    cp.keywords, cp.language, cp.comments = L["mots_cles"], L["langue_word"], ""
    cp.last_modified_by, cp.category = L["auteur"], L.get("categorie", "")
    return doc


# ── les sections ───────────────────────────────────────────────────────────
DEBUTS = {"suivante": WD_SECTION.NEW_PAGE, "impaire": WD_SECTION.ODD_PAGE, "paire": WD_SECTION.EVEN_PAGE,
          "continue": WD_SECTION.CONTINUOUS}


def regler_section(section, genre, debut=None, colonnes=1, premiere_page=True):
    """Marges, début et colonnes d'une section ; genre : couverture, liminaire, texte ou reference."""
    section.page_width, section.page_height = Cm(PAGE_L), Cm(PAGE_H)
    section.top_margin, section.bottom_margin = Cm(HAUT), Cm(BAS)
    section.header_distance, section.footer_distance = Cm(EN_TETE), Cm(PIED)
    section.gutter = Cm(0)
    if genre in ("reference", "couverture"):
        section.left_margin = section.right_margin = Cm(REFERENCE_MARGE)
    else:
        section.left_margin, section.right_margin = Cm(INTERIEURE), Cm(EXTERIEURE)
    if debut:
        section.start_type = DEBUTS[debut]
    section.different_first_page_header_footer = premiere_page
    spr = section._sectPr
    cols = X.xml(f'<w:cols w:num="{colonnes}" w:space="{X.twips(0.7)}"/>' if colonnes > 1
                 else f'<w:cols w:space="{X.twips(0.7)}"/>')
    X.poser(spr, cols)


def _vider(partie):
    """Vide un en-tête ou un pied (garde un paragraphe vide, sans le style « Header » de Word)."""
    partie.is_linked_to_previous = False
    for p in list(partie.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p = partie.paragraphs[0]
    for r in list(p._p):
        if r.tag != qn("w:pPr"):
            p._p.remove(r)
    return p


def _style(paragraphe, doc, nom):
    paragraphe.style = doc.styles[nom]


def entetes(section, doc, genre, langue):
    """Les en-têtes et les pieds d'une section, selon son genre (voir l'en-tête du module)."""
    L = langue.L
    rpr_titre = X.rpr_xml(police=TITRE, couleur=PRIMAIRE, taille=7.5)
    rpr_texte = X.rpr_xml(police=SANS, couleur=ARDOISE, taille=7.5)
    impair, pair = section.header, section.even_page_header
    pied_impair, pied_pair = section.footer, section.even_page_footer
    premier, pied_premier = section.first_page_header, section.first_page_footer
    for partie in (impair, pair, pied_impair, pied_pair, premier, pied_premier):
        _vider(partie)
    for partie in (premier, pied_premier):
        _style(partie.paragraphs[0], doc, "Poly En-tête")
    if genre == "couverture":
        for partie in (impair, pair, pied_impair, pied_pair):
            _style(partie.paragraphs[0], doc, "Poly En-tête")
        return
    debord = DEBORD if genre in ("texte", "liminaire") else 0.0
    # Les pieds : le numéro de page au coin extérieur.
    for partie, cote in ((pied_impair, "right"), (pied_pair, "left")):
        p = partie.paragraphs[0]
        _style(p, doc, "Poly Pied")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if cote == "right" else WD_ALIGN_PARAGRAPH.LEFT
        if cote == "right":
            p.paragraph_format.right_indent = Cm(-debord)
        else:
            p.paragraph_format.left_indent = Cm(-debord)
        X.ajouter_champ(p, "PAGE", "1", X.rpr_xml(police=TITRE, couleur=PRIMAIRE, taille=9))
    # Les en-têtes : pages paires, le titre du livre ; pages impaires, le chapitre ou la partie.
    for partie, cote in ((impair, "right"), (pair, "left")):
        p = partie.paragraphs[0]
        _style(p, doc, "Poly En-tête")
        if genre == "liminaire":
            continue
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if cote == "right" else WD_ALIGN_PARAGRAPH.LEFT
        if cote == "right":
            p.paragraph_format.right_indent = Cm(-debord)
        else:
            p.paragraph_format.left_indent = Cm(-debord)
        X.bordures_paragraphe(p, bottom=X.filet(sz=4, space=3, color=PRIMAIRE))
        if cote == "left":
            titre = L["titre"].replace("'", "’")
            p._p.append(X.xml(X.run_xml(titre, X.rpr_xml(police=SANS, couleur=ARDOISE, taille=7.5, espacement=0.3))))
        elif genre == "texte":
            X.ajouter_champ(p, 'STYLEREF "Poly Numéro de chapitre"', L["chapitre"].format(n=""), rpr_titre)
            p._p.append(X.xml(X.run_xml("  ·  ", rpr_texte)))
            X.ajouter_champ(p, 'STYLEREF "Poly Titre de chapitre"', "", rpr_texte)
        else:
            X.ajouter_champ(p, 'STYLEREF "Poly Titre de partie"', "", rpr_titre)


def nouvelle_section(doc, genre, langue, debut="impaire", colonnes=1, premiere_page=True):
    section = doc.add_section(DEBUTS[debut])
    regler_section(section, genre, debut, colonnes, premiere_page)
    entetes(section, doc, genre, langue)
    return section


def section_continue(doc, colonnes):
    """Une section continue qui change le nombre de colonnes ; elle garde les en-têtes de la précédente."""
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    spr = section._sectPr
    X.poser(spr, X.xml(f'<w:cols w:num="{colonnes}" w:space="{X.twips(0.7)}"/>' if colonnes > 1
                       else f'<w:cols w:space="{X.twips(0.7)}"/>'))
    for partie in (section.header, section.even_page_header, section.footer, section.even_page_footer,
                   section.first_page_header, section.first_page_footer):
        partie.is_linked_to_previous = True
    section.different_first_page_header_footer = False
    return section


def chemin_relatif(chemin, racine):
    return os.path.relpath(chemin, racine).replace("\\", "/")
